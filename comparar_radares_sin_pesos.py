"""
Cálculo de radar sin pesos (promedio simple de indicadores transformers)
y métricas de clasificación contra Clasificacion_radar_oficial_promedio y Clasificacion_IDIC.

Clasificación: 3 clases (Bajo / Medio / Alto) por terciles.
Rankings: posición 1 = mayor valor de radar (más actividad mediática detectada).
"""
import argparse
import os
import shutil
from typing import Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import (
    accuracy_score,
    classification_report,
    cohen_kappa_score,
    confusion_matrix,
    precision_recall_fscore_support,
)


INDICADORES_36 = [
    'participacion_comunitaria', 'incentivos_economicos', 'fortalecimiento_institucional',
    'impactos_ambientales', 'conflictos_socioambientales', 'desplazamiento_forzado',
    'reasentamiento', 'protesta_social', 'amenaza_intimidacion', 'consulta_previa',
    'audiencia_publica', 'taller_participativo', 'conflicto_territorial', 'rechazo_proyecto',
    'deficit_derechos', 'denuncia_violacion', 'deficit_participacion_efectiva', 'ruptura_dialogo',
    'reivindicacion_territorial', 'exclusion_participacion', 'equidad_inclusion', 'grupos_etnicos',
    'movimientos_sociales', 'grupos_poblacionales_afectados', 'participacion_economica_local',
    'transparencia_contractual', 'zonas_proteccion_alimentaria', 'respeto_territorios',
    'presencia_grupos_armados', 'desaparicion_lideres', 'grupos_etnicos_entidades',
    'grupos_armados_entidades', 'organizaciones_entidades', 'lideres_entidades',
    'instituciones_entidades', 'actores_economicos_entidades',
]

ETIQUETAS_CLASES = ['Bajo', 'Medio', 'Alto']


def _normalizar_nombre_depto(nombre: str) -> str:
    """Mapea variantes de nombre de departamento al esquema del xlsx."""
    txt = str(nombre).strip()
    equivalencias = {
        'San Andrés y Providencia': 'San Andrés',
        'San Andres y Providencia': 'San Andrés',
        'San Andres': 'San Andrés',
    }
    return equivalencias.get(txt, txt)


def _minmax_0_1(serie: pd.Series) -> pd.Series:
    smin, smax = serie.min(), serie.max()
    if smax > smin:
        return (serie - smin) / (smax - smin)
    return pd.Series([0.5] * len(serie), index=serie.index)


def _terciles(serie: pd.Series, invertir: bool = False) -> pd.Series:
    """
    Clasifica en 3 clases por terciles.
    invertir=True: valor alto → 'Bajo' (usado para radar_oficial donde alto = más riesgo).
    invertir=False: valor alto → 'Alto' (IDIC y radares nuevos).
    """
    if invertir:
        labels = ['Alto', 'Medio', 'Bajo']   # el tercil más bajo del valor → 'Alto' (menos riesgo)
    else:
        labels = ['Bajo', 'Medio', 'Alto']

    return pd.qcut(serie, q=3, labels=labels, duplicates='drop').astype(str)


def calcular_radar_sin_pesos(df_indicadores: pd.DataFrame, etiqueta: str) -> pd.DataFrame:
    """
    Promedio simple de los 36 indicadores → normalización 0-1 → terciles → ranking.
    Ranking: posición 1 = mayor valor (más actividad mediática detectada).
    """
    if 'departamento' not in df_indicadores.columns:
        raise ValueError(f"[{etiqueta}] CSV debe contener columna 'departamento'")

    df = df_indicadores.copy()
    df['departamento'] = df['departamento'].astype(str).str.strip().map(_normalizar_nombre_depto)

    faltantes = [c for c in INDICADORES_36 if c not in df.columns]
    if faltantes:
        print(f"[{etiqueta}] Indicadores faltantes (se rellenan con 0): {faltantes}")
        for c in faltantes:
            df[c] = 0.0

    for c in INDICADORES_36:
        df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0.0)

    radar_raw = df[INDICADORES_36].mean(axis=1)
    radar_norm = _minmax_0_1(radar_raw)
    clasif = _terciles(radar_norm, invertir=False)   # alto radar = 'Alto'
    ranking = radar_norm.rank(method='min', ascending=False).astype(int)

    out = pd.DataFrame({
        'departamento': df['departamento'].values,
        f'radar_{etiqueta}_promedio_normalizado': radar_norm.values,
        f'Clasificacion_radar_{etiqueta}': clasif.values,
        f'Ranking_radar_{etiqueta}': ranking.values,
    })
    return out


def _reclasificar_referencias_en_xlsx(ws, headers_existentes: dict) -> None:
    """
    Reclasifica Clasificacion_radar_oficial_promedio y Clasificacion_IDIC
    desde sus valores numéricos originales a 3 clases por terciles.
    """
    col_depto = headers_existentes.get('Departamento') or headers_existentes.get('departamento')
    col_radar_val = headers_existentes.get('radar_oficial_promedio')
    col_idic_val = headers_existentes.get('IDIC')
    col_clasif_radar = headers_existentes.get('Clasificacion_radar_oficial_promedio')
    col_clasif_idic = headers_existentes.get('Clasificacion_IDIC')

    if not all([col_radar_val, col_idic_val, col_clasif_radar, col_clasif_idic]):
        print("ADVERTENCIA: no se encontraron todas las columnas de referencia para reclasificar.")
        return

    # Leer valores numéricos
    filas = list(range(2, ws.max_row + 1))
    radar_vals = []
    idic_vals = []
    for r in filas:
        rv = ws.cell(r, col_radar_val).value
        iv = ws.cell(r, col_idic_val).value
        radar_vals.append(float(rv) if rv is not None else np.nan)
        idic_vals.append(float(iv) if iv is not None else np.nan)

    s_radar = pd.Series(radar_vals)
    s_idic  = pd.Series(idic_vals)

    # radar_oficial: alto valor = más riesgo = peor → invertir
    clasif_radar = _terciles(s_radar, invertir=True)
    # IDIC: alto valor = mejor situación → no invertir
    clasif_idic  = _terciles(s_idic, invertir=False)

    for i, r in enumerate(filas):
        ws.cell(r, col_clasif_radar, clasif_radar.iloc[i])
        ws.cell(r, col_clasif_idic,  clasif_idic.iloc[i])

    print("  Clasificacion_radar_oficial_promedio -> reclasificada (3 clases, invertida)")
    print("  Clasificacion_IDIC                  -> reclasificada (3 clases)")


def actualizar_xlsx(ruta_xlsx: str, df_viejo: pd.DataFrame, df_nuevo: pd.DataFrame) -> None:
    """
    Actualiza el xlsx:
    - Reclasifica las referencias a 3 clases.
    - Agrega/actualiza columnas de radar viejo y nuevo (valor, clasificación, ranking).
    """
    if not os.path.isfile(ruta_xlsx):
        raise FileNotFoundError(f"xlsx no existe: {ruta_xlsx}")

    backup = ruta_xlsx.replace('.xlsx', '_backup.xlsx')
    if not os.path.isfile(backup):
        shutil.copy2(ruta_xlsx, backup)
        print(f"Backup creado: {backup}")

    wb = load_workbook(ruta_xlsx)
    ws = wb[wb.sheetnames[0]]

    headers_existentes = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if h is not None:
            headers_existentes[str(h).strip()] = c

    # 1. Reclasificar referencias
    _reclasificar_referencias_en_xlsx(ws, headers_existentes)

    # 2. Preparar columnas de los nuevos radares (6 columnas: valor + clasif + ranking × 2)
    nuevas_cols = [
        'radar_viejo_promedio_normalizado',
        'Clasificacion_radar_viejo',
        'Ranking_radar_viejo',
        'radar_nuevo_promedio_normalizado',
        'Clasificacion_radar_nuevo',
        'Ranking_radar_nuevo',
    ]
    col_indices = {}
    siguiente_col = ws.max_column + 1
    for nombre in nuevas_cols:
        if nombre in headers_existentes:
            col_indices[nombre] = headers_existentes[nombre]
        else:
            ws.cell(row=1, column=siguiente_col, value=nombre)
            col_indices[nombre] = siguiente_col
            siguiente_col += 1

    col_depto = headers_existentes.get('Departamento') or headers_existentes.get('departamento')
    if col_depto is None:
        raise ValueError("xlsx no contiene columna 'Departamento'")

    mapa_viejo = {
        _normalizar_nombre_depto(r['departamento']): (
            r['radar_viejo_promedio_normalizado'],
            r['Clasificacion_radar_viejo'],
            r['Ranking_radar_viejo'],
        )
        for _, r in df_viejo.iterrows()
    }
    mapa_nuevo = {
        _normalizar_nombre_depto(r['departamento']): (
            r['radar_nuevo_promedio_normalizado'],
            r['Clasificacion_radar_nuevo'],
            r['Ranking_radar_nuevo'],
        )
        for _, r in df_nuevo.iterrows()
    }

    no_match = []
    for r in range(2, ws.max_row + 1):
        dep = ws.cell(row=r, column=col_depto).value
        if dep is None:
            continue
        dep_norm = _normalizar_nombre_depto(dep)

        if dep_norm in mapa_viejo:
            val, cls, rank = mapa_viejo[dep_norm]
            ws.cell(row=r, column=col_indices['radar_viejo_promedio_normalizado'], value=float(val))
            ws.cell(row=r, column=col_indices['Clasificacion_radar_viejo'],        value=str(cls))
            ws.cell(row=r, column=col_indices['Ranking_radar_viejo'],              value=int(rank))
        else:
            no_match.append(('viejo', dep_norm))

        if dep_norm in mapa_nuevo:
            val, cls, rank = mapa_nuevo[dep_norm]
            ws.cell(row=r, column=col_indices['radar_nuevo_promedio_normalizado'], value=float(val))
            ws.cell(row=r, column=col_indices['Clasificacion_radar_nuevo'],        value=str(cls))
            ws.cell(row=r, column=col_indices['Ranking_radar_nuevo'],              value=int(rank))
        else:
            no_match.append(('nuevo', dep_norm))

    if no_match:
        print(f"ADVERTENCIA: departamentos sin match: {no_match}")

    wb.save(ruta_xlsx)
    print(f"xlsx actualizado: {ruta_xlsx}")


def _calcular_metricas_par(y_true: pd.Series, y_pred: pd.Series, etiqueta: str) -> Tuple[dict, pd.DataFrame, np.ndarray]:
    """Calcula métricas de clasificación para un par (verdadero, predicho)."""
    mask = (y_true.notna() & y_pred.notna() &
            (y_true.astype(str).str.strip() != '') &
            (y_pred.astype(str).str.strip() != ''))
    yt = y_true[mask].astype(str).tolist()
    yp = y_pred[mask].astype(str).tolist()

    acc = accuracy_score(yt, yp)
    p_macro, r_macro, f_macro, _ = precision_recall_fscore_support(
        yt, yp, labels=ETIQUETAS_CLASES, average='macro', zero_division=0
    )
    kappa = cohen_kappa_score(yt, yp, labels=ETIQUETAS_CLASES)

    resumen = {
        'comparacion': etiqueta,
        'n_departamentos': len(yt),
        'accuracy': float(acc),
        'precision_macro': float(p_macro),
        'recall_macro': float(r_macro),
        'f1_macro': float(f_macro),
        'cohen_kappa': float(kappa),
    }

    rep = classification_report(
        yt, yp, labels=ETIQUETAS_CLASES, target_names=ETIQUETAS_CLASES,
        output_dict=True, zero_division=0,
    )
    detalle_rows = []
    for clase in ETIQUETAS_CLASES:
        d = rep[clase]
        detalle_rows.append({
            'comparacion': etiqueta,
            'clase': clase,
            'precision': float(d['precision']),
            'recall': float(d['recall']),
            'f1-score': float(d['f1-score']),
            'support': int(d['support']),
        })
    detalle_df = pd.DataFrame(detalle_rows)

    cm = confusion_matrix(yt, yp, labels=ETIQUETAS_CLASES)
    return resumen, detalle_df, cm


def metricas_clasificacion(ruta_xlsx: str, ruta_salida: str) -> None:
    """Genera xlsx con métricas de clasificación para las 4 comparaciones."""
    wb = load_workbook(ruta_xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    data = {}
    for c_idx, h in enumerate(headers, start=1):
        if h is None:
            continue
        col = [ws.cell(row=r, column=c_idx).value for r in range(2, ws.max_row + 1)]
        data[str(h).strip()] = col
    df = pd.DataFrame(data)

    pares = [
        ('viejo_vs_radar_oficial', 'Clasificacion_radar_viejo',  'Clasificacion_radar_oficial_promedio'),
        ('viejo_vs_IDIC',          'Clasificacion_radar_viejo',  'Clasificacion_IDIC'),
        ('nuevo_vs_radar_oficial', 'Clasificacion_radar_nuevo',  'Clasificacion_radar_oficial_promedio'),
        ('nuevo_vs_IDIC',          'Clasificacion_radar_nuevo',  'Clasificacion_IDIC'),
    ]

    for etiqueta, col_pred, col_true in pares:
        for c in (col_pred, col_true):
            if c not in df.columns:
                raise ValueError(f"Falta columna '{c}' en xlsx")

    resumenes, detalles, matrices = [], [], {}

    for etiqueta, col_pred, col_true in pares:
        resumen, detalle, cm = _calcular_metricas_par(df[col_true], df[col_pred], etiqueta)
        resumenes.append(resumen)
        detalles.append(detalle)
        matrices[etiqueta] = cm

    df_resumen = pd.DataFrame(resumenes)
    df_detalle = pd.concat(detalles, ignore_index=True)

    filas_cm = []
    for etiqueta, cm in matrices.items():
        filas_cm.append({'comparacion': etiqueta, 'verdadero \\ predicho': '---',
                         **{c: '' for c in ETIQUETAS_CLASES}})
        for i, clase_real in enumerate(ETIQUETAS_CLASES):
            fila = {'comparacion': etiqueta, 'verdadero \\ predicho': clase_real}
            for j, clase_pred in enumerate(ETIQUETAS_CLASES):
                fila[clase_pred] = int(cm[i, j])
            filas_cm.append(fila)
        filas_cm.append({'comparacion': '', 'verdadero \\ predicho': ''})
    df_cm = pd.DataFrame(filas_cm)

    with pd.ExcelWriter(ruta_salida, engine='openpyxl') as writer:
        df_resumen.to_excel(writer, sheet_name='resumen', index=False)
        df_detalle.to_excel(writer, sheet_name='detalle_por_clase', index=False)
        df_cm.to_excel(writer, sheet_name='matrices_confusion', index=False)

    print(f"\nMétricas exportadas: {ruta_salida}")
    print("\n=== RESUMEN ===")
    print(df_resumen.to_string(index=False))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--csv-viejo', required=True)
    parser.add_argument('--csv-nuevo', required=True)
    parser.add_argument('--xlsx-comparacion', required=True)
    parser.add_argument('--salida-metricas', required=True)
    args = parser.parse_args()

    df_viejo_csv = pd.read_csv(args.csv_viejo)
    df_nuevo_csv = pd.read_csv(args.csv_nuevo)
    print(f"CSV viejo: {df_viejo_csv.shape}")
    print(f"CSV nuevo: {df_nuevo_csv.shape}")

    df_viejo = calcular_radar_sin_pesos(df_viejo_csv, 'viejo')
    df_nuevo  = calcular_radar_sin_pesos(df_nuevo_csv,  'nuevo')

    print("\n=== Radar viejo (ordenado por ranking) ===")
    print(df_viejo.sort_values('Ranking_radar_viejo').to_string(index=False))
    print("\n=== Radar nuevo (ordenado por ranking) ===")
    print(df_nuevo.sort_values('Ranking_radar_nuevo').to_string(index=False))

    actualizar_xlsx(args.xlsx_comparacion, df_viejo, df_nuevo)
    metricas_clasificacion(args.xlsx_comparacion, args.salida_metricas)


if __name__ == '__main__':
    main()
