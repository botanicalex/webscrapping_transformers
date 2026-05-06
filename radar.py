import argparse
import json
import os
import random
import re
import shutil
from typing import Any, Dict, List, Optional, Sequence

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler

import config_pipeline as cfg
import metricas_y_calculo_de_error as mc


class CalculadorRadar:
    MAPEO_PERIODICO_DEPARTAMENTO = {
        'El Colombiano': 'Antioquia', 'El Diario': 'Risaralda', 'BC Noticias': 'Caldas', 'El Quindiano': 'Quindío',
        'El País Cali': 'Valle del Cauca', 'Diario Occidente': 'Valle del Cauca', 'Diario del Sur': 'Nariño',
        'Diario del Cauca': 'Cauca', 'Chocó 7 Días': 'Chocó', 'Llano al Mundo': 'Meta', 'Diario de Casanare': 'Casanare',
        'La Voz del Cinaruco': 'Arauca', 'El Morichal': 'Vichada', 'Mi Putumayo': 'Putumayo', 'El Tiempo': 'Cundinamarca',
        'La República': 'Cundinamarca', 'Portafolio': 'Cundinamarca', 'Publimetro': 'Cundinamarca', 'Las2Orillas': 'Cundinamarca',
        'El Heraldo': 'Atlántico', 'El Universal': 'Bolívar', 'El Pilón': 'Cesar', 'El Meridiano': 'Córdoba',
        'Vanguardia': 'Santander', 'Trochando Sin Fronteras': 'Arauca', 'Enlace Television': 'Santander', 'Corrillos': 'Santander'
    }

    BLOQUES_PCA = {
        'bloque_A': ['denuncia_violacion', 'transparencia_contractual', 'conflicto_territorial', 'actores_economicos_entidades'],
        'bloque_B': ['presencia_grupos_armados', 'desaparicion_lideres', 'amenaza_intimidacion', 'grupos_armados_entidades', 'conflictos_socioambientales'],
        'bloque_C': ['fortalecimiento_institucional', 'llamado_dialogo', 'instituciones_entidades', 'propuesta_alternativa', 'incentivos_economicos'],
        'bloque_D': ['desplazamiento_forzado', 'grupos_etnicos', 'grupos_poblacionales_afectados', 'zonas_proteccion_alimentaria', 'respeto_territorios', 'equidad_inclusion', 'grupos_etnicos_entidades'],
        'bloque_E': ['participacion_comunitaria', 'consulta_previa', 'audiencia_publica', 'taller_participativo', 'exigencia_participacion', 'movimientos_sociales', 'organizaciones_entidades', 'lideres_entidades', 'nivel_acuerdo_proyecto']
    }

    VARS_INVERTIR = {
        'fortalecimiento_institucional', 'llamado_dialogo', 'consulta_previa', 'audiencia_publica',
        'taller_participativo', 'participacion_comunitaria', 'nivel_acuerdo_proyecto', 'propuesta_alternativa',
        'incentivos_economicos', 'participacion_economica_local', 'instituciones_entidades'
    }

    COLUMNAS_BINARIAS = [
        'participacion_comunitaria', 'incentivos_economicos', 'fortalecimiento_institucional', 'impactos_ambientales', 'conflictos_socioambientales',
        'desplazamiento_forzado', 'reasentamiento', 'protesta_social', 'amenaza_intimidacion', 'consulta_previa', 'audiencia_publica', 'taller_participativo', 'conflicto_territorial',
        'nivel_acuerdo_proyecto', 'demanda_derechos', 'denuncia_violacion', 'propuesta_alternativa', 'llamado_dialogo', 'defensa_territorio', 'exigencia_participacion',
        'equidad_inclusion', 'grupos_etnicos', 'movimientos_sociales', 'grupos_poblacionales_afectados', 'participacion_economica_local', 'transparencia_contractual',
        'zonas_proteccion_alimentaria', 'respeto_territorios', 'presencia_grupos_armados', 'desaparicion_lideres',
        'grupos_etnicos_entidades', 'grupos_armados_entidades', 'organizaciones_entidades', 'lideres_entidades', 'instituciones_entidades', 'actores_economicos_entidades'
    ]

    OPERACION_BLOQUES = "bloques"
    OPERACION_INDICADORES = "indicadores_transformers"
    OPERACIONES = (OPERACION_BLOQUES, OPERACION_INDICADORES)

    PESOS_BLOQUES_FIJOS = {
        "corrupcion_bloque_B": 0.80,
        "corrupcion_bloque_A": 0.35,
        "corrupcion_bloque_C": -0.15,
        "vulneracion_bloque_D": 0.50,
        "vulneracion_bloque_E": 0.50,
        "radar_corrupcion": 0.60,
        "radar_vulneracion": 0.40,
    }

    PESOS_INDICADORES_FIJOS = {
        "bloque_A": 0.20,
        "bloque_B": 0.25,
        "bloque_C": 0.15,
        "bloque_D": 0.20,
        "bloque_E": 0.20,
        "corrupcion": 0.60,
        "vulneracion": 0.40,
    }

    def calcular(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        if 'departamento' not in df.columns or not df['departamento'].notna().any():
            df['departamento'] = df['periodico'].map(self.MAPEO_PERIODICO_DEPARTAMENTO)
            df = df.dropna(subset=['departamento'])
        cols = [c for c in self.COLUMNAS_BINARIAS if c in df.columns]
        for c in cols:
            if df[c].dtype == bool:
                df[c] = df[c].astype(float)
            else:
                df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0.0)
        df_tasas = df.groupby('departamento')[cols].mean()
        df_tasas['n_articulos'] = df.groupby('departamento').size()
        for c in [c for c in cols if c in self.VARS_INVERTIR]:
            df_tasas[c] = 1 - df_tasas[c]

        df_bloques = pd.DataFrame(index=df_tasas.index)
        n_depts = len(df_tasas)
        for b, variables in self.BLOQUES_PCA.items():
            v = [x for x in variables if x in df_tasas.columns]
            if not v:
                df_bloques[b] = 0.0
                continue
            X = df_tasas[v].values
            if n_depts < len(v):
                scores = X.mean(axis=1)
            else:
                try:
                    Xs = StandardScaler().fit_transform(X)
                    scores = PCA(n_components=1).fit_transform(Xs)[:, 0]
                except Exception:
                    scores = X.mean(axis=1)
            smin, smax = scores.min(), scores.max()
            if smax > smin:
                sn = (scores - smin) / (smax - smin) * 100
            elif n_depts == 1:
                sn = np.clip(scores, 0, 1) * 100
            else:
                sn = np.full_like(scores, 50.0, dtype=float)
            df_bloques[b] = sn

        df_bloques['n_articulos'] = df_tasas['n_articulos']
        df_sub = df_bloques.copy()
        corr_raw = 0.80 * df_sub['bloque_B'] + 0.35 * df_sub['bloque_A'] - 0.15 * df_sub['bloque_C']
        vul_raw = 0.50 * df_sub['bloque_D'] + 0.50 * df_sub['bloque_E']
        df_sub['corrupcion_score'] = self._minmax(corr_raw)
        df_sub['vulneracion_score'] = self._minmax(vul_raw)
        radar_raw = 0.60 * df_sub['corrupcion_score'] + 0.40 * df_sub['vulneracion_score']
        df_sub['radar_propio'] = self._minmax(radar_raw)
        p25 = df_sub['radar_propio'].quantile(0.25)
        p75 = df_sub['radar_propio'].quantile(0.75)

        def cat(x: float) -> str:
            if x <= p25:
                return "bajo"
            if x >= p75:
                return "alto"
            return "medio"

        df_sub['categoria_riesgo'] = df_sub['radar_propio'].apply(cat)
        out = df_sub.reset_index().rename(columns={'index': 'departamento'})
        return out[['departamento', 'n_articulos', 'bloque_A', 'bloque_B', 'bloque_C', 'bloque_D', 'bloque_E', 'corrupcion_score', 'vulneracion_score', 'radar_propio', 'categoria_riesgo']].sort_values('radar_propio', ascending=False)

    def calcular_desde_indicadores(
        self,
        df_indicadores: pd.DataFrame,
        df_radar_base: Optional[pd.DataFrame] = None,
        operacion: str = OPERACION_BLOQUES,
        pesos: Optional[Dict[str, float]] = None,
    ) -> pd.DataFrame:
        df_tasas = self._preparar_tasas_indicadores(df_indicadores)
        n_articulos = self._resolver_n_articulos(df_radar_base, df_tasas.index)
        if operacion == self.OPERACION_BLOQUES:
            return self._calcular_bloques_desde_tasas(df_tasas, n_articulos)
        if operacion == self.OPERACION_INDICADORES:
            return self._calcular_indicadores_transformers(df_tasas, n_articulos, pesos)
        raise ValueError(f"Operación no soportada: '{operacion}'")

    def _preparar_tasas_indicadores(self, df_indicadores: pd.DataFrame) -> pd.DataFrame:
        if 'departamento' not in df_indicadores.columns:
            raise ValueError("El archivo de indicadores debe contener la columna 'departamento'")
        df = df_indicadores.copy()
        df['departamento'] = df['departamento'].astype(str).str.strip()
        df = df[df['departamento'] != ""]
        for c in self.COLUMNAS_BINARIAS:
            if c not in df.columns:
                df[c] = np.nan
            df[c] = pd.to_numeric(df[c], errors='coerce')
        return df.groupby('departamento')[self.COLUMNAS_BINARIAS].mean().fillna(0.0)

    def _resolver_n_articulos(self, df_radar_base: Optional[pd.DataFrame], index_departamentos: pd.Index) -> pd.Series:
        n_articulos = pd.Series([np.nan] * len(index_departamentos), index=index_departamentos, dtype=float)
        if df_radar_base is None or df_radar_base.empty:
            return n_articulos
        if 'departamento' not in df_radar_base.columns or 'n_articulos' not in df_radar_base.columns:
            return n_articulos
        base = df_radar_base[['departamento', 'n_articulos']].copy()
        base['departamento'] = base['departamento'].astype(str).str.strip()
        base['n_articulos'] = pd.to_numeric(base['n_articulos'], errors='coerce')
        mapa = base.dropna(subset=['departamento']).drop_duplicates(subset=['departamento']).set_index('departamento')['n_articulos']
        valores = index_departamentos.to_series().map(mapa)
        return valores.astype(float)

    def _calcular_bloques_desde_tasas(self, df_tasas: pd.DataFrame, n_articulos: pd.Series) -> pd.DataFrame:
        df_tasas = df_tasas.copy()
        cols = [c for c in self.COLUMNAS_BINARIAS if c in df_tasas.columns]
        for c in [c for c in cols if c in self.VARS_INVERTIR]:
            df_tasas[c] = 1 - df_tasas[c]

        df_bloques = pd.DataFrame(index=df_tasas.index)
        n_depts = len(df_tasas)
        for b, variables in self.BLOQUES_PCA.items():
            v = [x for x in variables if x in df_tasas.columns]
            if not v:
                df_bloques[b] = 0.0
                continue
            X = df_tasas[v].values
            if n_depts < len(v):
                scores = X.mean(axis=1)
            else:
                try:
                    Xs = StandardScaler().fit_transform(X)
                    scores = PCA(n_components=1).fit_transform(Xs)[:, 0]
                except Exception:
                    scores = X.mean(axis=1)
            smin, smax = scores.min(), scores.max()
            if smax > smin:
                sn = (scores - smin) / (smax - smin) * 100
            elif n_depts == 1:
                sn = np.clip(scores, 0, 1) * 100
            else:
                sn = np.full_like(scores, 50.0, dtype=float)
            df_bloques[b] = sn

        df_bloques['n_articulos'] = n_articulos.reindex(df_bloques.index)
        df_sub = df_bloques.copy()
        corr_raw = 0.80 * df_sub['bloque_B'] + 0.35 * df_sub['bloque_A'] - 0.15 * df_sub['bloque_C']
        vul_raw = 0.50 * df_sub['bloque_D'] + 0.50 * df_sub['bloque_E']
        df_sub['corrupcion_score'] = self._minmax(corr_raw)
        df_sub['vulneracion_score'] = self._minmax(vul_raw)
        radar_raw = 0.60 * df_sub['corrupcion_score'] + 0.40 * df_sub['vulneracion_score']
        df_sub['radar_propio'] = self._minmax(radar_raw)
        p25 = df_sub['radar_propio'].quantile(0.25)
        p75 = df_sub['radar_propio'].quantile(0.75)

        def cat(x: float) -> str:
            if x <= p25:
                return "bajo"
            if x >= p75:
                return "alto"
            return "medio"

        df_sub['categoria_riesgo'] = df_sub['radar_propio'].apply(cat)
        out = df_sub.reset_index().rename(columns={'index': 'departamento'})
        return out[['departamento', 'n_articulos', 'bloque_A', 'bloque_B', 'bloque_C', 'bloque_D', 'bloque_E', 'corrupcion_score', 'vulneracion_score', 'radar_propio', 'categoria_riesgo']].sort_values('radar_propio', ascending=False)

    def _calcular_indicadores_transformers(
        self,
        df_tasas: pd.DataFrame,
        n_articulos: pd.Series,
        pesos: Optional[Dict[str, float]],
    ) -> pd.DataFrame:
        df_riesgo = df_tasas.copy()
        for c in [c for c in self.COLUMNAS_BINARIAS if c in self.VARS_INVERTIR and c in df_riesgo.columns]:
            df_riesgo[c] = 1 - df_riesgo[c]

        df_bloques = pd.DataFrame(index=df_riesgo.index)
        for b, variables in self.BLOQUES_PCA.items():
            v = [x for x in variables if x in df_riesgo.columns]
            if not v:
                df_bloques[b] = 0.0
            else:
                df_bloques[b] = df_riesgo[v].mean(axis=1) * 100.0

        pesos_finales = self.normalizar_pesos_indicadores(pesos)
        corr_raw = (
            pesos_finales['bloque_B'] * df_bloques['bloque_B']
            + pesos_finales['bloque_A'] * df_bloques['bloque_A']
            + pesos_finales['bloque_C'] * df_bloques['bloque_C']
        )
        vul_raw = (
            pesos_finales['bloque_D'] * df_bloques['bloque_D']
            + pesos_finales['bloque_E'] * df_bloques['bloque_E']
        )

        df_sub = df_bloques.copy()
        df_sub['n_articulos'] = n_articulos.reindex(df_sub.index)
        df_sub['corrupcion_score'] = self._minmax(corr_raw)
        df_sub['vulneracion_score'] = self._minmax(vul_raw)
        radar_raw = (
            pesos_finales['corrupcion'] * df_sub['corrupcion_score']
            + pesos_finales['vulneracion'] * df_sub['vulneracion_score']
        )
        df_sub['radar_propio'] = self._minmax(radar_raw)

        p25 = df_sub['radar_propio'].quantile(0.25)
        p75 = df_sub['radar_propio'].quantile(0.75)

        def cat(x: float) -> str:
            if x <= p25:
                return "bajo"
            if x >= p75:
                return "alto"
            return "medio"

        df_sub['categoria_riesgo'] = df_sub['radar_propio'].apply(cat)
        out = df_sub.reset_index().rename(columns={'index': 'departamento'})
        return out[['departamento', 'n_articulos', 'bloque_A', 'bloque_B', 'bloque_C', 'bloque_D', 'bloque_E', 'corrupcion_score', 'vulneracion_score', 'radar_propio', 'categoria_riesgo']].sort_values('radar_propio', ascending=False)

    @classmethod
    def normalizar_pesos_indicadores(cls, pesos: Optional[Dict[str, float]]) -> Dict[str, float]:
        fuente = dict(cls.PESOS_INDICADORES_FIJOS)
        if isinstance(pesos, dict):
            for k, v in pesos.items():
                fuente[k] = v
        bloques = cls._normalizar_pesos(fuente, ['bloque_A', 'bloque_B', 'bloque_C', 'bloque_D', 'bloque_E'])
        radar = cls._normalizar_pesos(fuente, ['corrupcion', 'vulneracion'])
        return {**bloques, **radar}

    @staticmethod
    def _normalizar_pesos(pesos: Dict[str, Any], claves: Sequence[str]) -> Dict[str, float]:
        limpios: Dict[str, float] = {}
        for c in claves:
            valor = pesos.get(c, 0.0)
            try:
                numero = float(valor)
            except (TypeError, ValueError):
                numero = 0.0
            if not np.isfinite(numero) or numero < 0:
                numero = 0.0
            limpios[c] = numero
        total = float(sum(limpios.values()))
        if total <= 0:
            uniforme = 1.0 / float(len(claves))
            return {c: uniforme for c in claves}
        return {c: limpios[c] / total for c in claves}

    @staticmethod
    def _minmax(s: pd.Series) -> pd.Series:
        smin, smax = s.min(), s.max()
        if smax > smin:
            return (s - smin) / (smax - smin) * 100
        return pd.Series([50.0] * len(s), index=s.index)


def _es_celda_vacia(valor: Any) -> bool:
    if valor is None:
        return True
    if isinstance(valor, float) and pd.isna(valor):
        return True
    return str(valor).strip() == ""


def _detectar_columna_vacia_en_hoja(ws: Worksheet) -> Optional[int]:
    max_col = ws.max_column
    max_row = ws.max_row
    for col_idx in range(1, max_col + 1):
        header = ws.cell(row=1, column=col_idx).value
        header_txt = "" if header is None else str(header).strip()
        if header_txt != "" and not header_txt.lower().startswith("unnamed:"):
            continue
        if max_row <= 1:
            return col_idx
        columna_vacia = True
        for row_idx in range(2, max_row + 1):
            if not _es_celda_vacia(ws.cell(row=row_idx, column=col_idx).value):
                columna_vacia = False
                break
        if columna_vacia:
            return col_idx
    return None


def _asegurar_excel_base(excel: str, departamentos: pd.Series) -> str:
    if os.path.isfile(excel):
        wb = load_workbook(excel)
        return wb.sheetnames[0]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="departamento")
    ws.cell(row=1, column=2, value="radar_oficial_promedio")
    dep_vals = departamentos.dropna().astype(str).str.strip().drop_duplicates().tolist()
    for i, dep in enumerate(dep_vals, start=2):
        ws.cell(row=i, column=1, value=dep)
        ws.cell(row=i, column=2, value="None")
    wb.save(excel)
    return ws.title


def _max_indice_experimento_en_excel(excel: str) -> int:
    if not os.path.isfile(excel):
        return 0
    wb = load_workbook(excel, data_only=True)
    ws = wb[wb.sheetnames[0]]
    patron = re.compile(r"^experimento_(\d+)$", flags=re.IGNORECASE)
    max_idx = 0
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        texto = str(header).strip()
        m = patron.match(texto)
        if m:
            max_idx = max(max_idx, int(m.group(1)))
    return max_idx


def _max_indice_experimento_en_log(archivo_log: str) -> int:
    if not os.path.isfile(archivo_log):
        return 0
    patron = re.compile(r"^experimento_(\d+)$", flags=re.IGNORECASE)
    max_idx = 0
    with open(archivo_log, "r", encoding="utf-8") as f:
        for linea in f:
            texto = linea.strip()
            if not texto:
                continue
            try:
                payload = json.loads(texto)
            except json.JSONDecodeError:
                continue
            nombre = str(payload.get("EXPERIMENTO_N", "")).strip()
            m = patron.match(nombre)
            if m:
                max_idx = max(max_idx, int(m.group(1)))
    return max_idx


def _siguiente_indice_experimento(excel: str, archivo_log: str) -> int:
    return max(_max_indice_experimento_en_excel(excel), _max_indice_experimento_en_log(archivo_log)) + 1


def _parsear_indice_experimento(nombre_experimento: str) -> Optional[int]:
    patron = re.compile(r"^experimento_(\d+)$", flags=re.IGNORECASE)
    match = patron.match(str(nombre_experimento).strip())
    if not match:
        return None
    return int(match.group(1))


def _actualizar_excel_experimento(excel: str, df_radar: pd.DataFrame, nombre_experimento: str) -> None:
    if 'departamento' not in df_radar.columns or 'radar_propio' not in df_radar.columns:
        raise ValueError("El resultado de radar debe contener 'departamento' y 'radar_propio'")

    df_radar_local = df_radar[['departamento', 'radar_propio']].copy()
    df_radar_local['departamento'] = df_radar_local['departamento'].astype(str).str.strip()
    df_radar_local['radar_propio'] = pd.to_numeric(df_radar_local['radar_propio'], errors='coerce')

    sheet_name = _asegurar_excel_base(excel, df_radar_local['departamento'])
    with pd.ExcelFile(excel, engine="openpyxl") as xl:
        df_excel = xl.parse(sheet_name)

    if 'departamento' not in df_excel.columns:
        raise ValueError("El Excel debe contener la columna 'departamento'")

    df_excel['departamento'] = df_excel['departamento'].astype(str).str.strip()
    if nombre_experimento in df_excel.columns:
        raise ValueError(f"La columna '{nombre_experimento}' ya existe en '{excel}'")

    wb = load_workbook(excel)
    ws = wb[sheet_name]
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        if str(header).strip() == nombre_experimento:
            raise ValueError(f"La columna '{nombre_experimento}' ya existe en '{excel}'")

    idx_departamento = {
        str(ws.cell(row=row_idx, column=1).value).strip().casefold(): row_idx
        for row_idx in range(2, ws.max_row + 1)
    }
    for dep in df_radar_local['departamento'].dropna().astype(str).str.strip().tolist():
        clave = dep.casefold()
        if clave not in idx_departamento:
            nueva_fila = ws.max_row + 1
            ws.cell(row=nueva_fila, column=1, value=dep)
            if ws.max_column >= 2 and _es_celda_vacia(ws.cell(row=nueva_fila, column=2).value):
                ws.cell(row=nueva_fila, column=2, value="None")
            idx_departamento[clave] = nueva_fila

    col_objetivo = _detectar_columna_vacia_en_hoja(ws)
    if col_objetivo is None:
        col_objetivo = ws.max_column + 1
    ws.cell(row=1, column=col_objetivo, value=nombre_experimento)

    radar_map = {
        k.casefold(): v
        for k, v in zip(df_radar_local['departamento'], df_radar_local['radar_propio'])
    }
    for row_idx in range(2, ws.max_row + 1):
        dep = ws.cell(row=row_idx, column=1).value
        clave = str(dep).strip().casefold()
        valor = radar_map.get(clave)
        if pd.isna(valor):
            ws.cell(row=row_idx, column=col_objetivo, value="None")
        else:
            ws.cell(row=row_idx, column=col_objetivo, value=float(valor))

    wb.save(excel)


def _validar_consistencia(df_indicadores: pd.DataFrame, df_resultado: pd.DataFrame) -> None:
    if df_resultado.empty:
        raise ValueError("El resultado del radar está vacío")
    if 'departamento' not in df_indicadores.columns or 'departamento' not in df_resultado.columns:
        raise ValueError("Los datos de entrada y salida deben contener la columna 'departamento'")
    if 'radar_propio' not in df_resultado.columns:
        raise ValueError("La salida de radar debe contener la columna 'radar_propio'")
    deps_in = set(df_indicadores['departamento'].dropna().astype(str).str.strip().tolist())
    deps_out = set(df_resultado['departamento'].dropna().astype(str).str.strip().tolist())
    faltantes = sorted(list(deps_in - deps_out))
    if faltantes:
        raise ValueError(f"Faltan departamentos en la salida de radar: {faltantes[:10]}")
    radar_numerico = pd.to_numeric(df_resultado['radar_propio'], errors='coerce')
    if radar_numerico.isna().all():
        raise ValueError("La salida contiene radar_propio no numérico para todos los departamentos")


def _append_log_experimentos(archivo_log: str, payload: Dict[str, Any]) -> None:
    directorio = os.path.dirname(archivo_log)
    if directorio:
        os.makedirs(directorio, exist_ok=True)
    with open(archivo_log, "a", encoding="utf-8") as f:
        f.write(json.dumps(payload, ensure_ascii=False) + "\n")


def _normalizar_id_experimento(nombre: Any) -> str:
    texto = str(nombre).strip()
    m = re.match(r"^experimento_(\d+)$", texto, flags=re.IGNORECASE)
    if m:
        return f"EXPERIMENTO_{int(m.group(1))}"
    return texto


def _es_id_experimento(nombre: Any) -> bool:
    texto = _normalizar_id_experimento(nombre)
    return re.match(r"^EXPERIMENTO_\d+$", texto, flags=re.IGNORECASE) is not None


def _nombre_carpeta_seguro(nombre: str) -> str:
    nombre_base = str(nombre).strip()
    if not nombre_base:
        nombre_base = "experimento"
    nombre_limpio = "".join(
        c if c.isalnum() or c in (" ", "-", "_") else "_" for c in nombre_base
    ).strip()
    if not nombre_limpio:
        return "experimento"
    return nombre_limpio[:120]


def _listar_experimentos_en_excel(excel: str) -> List[str]:
    if not os.path.isfile(excel):
        return []
    wb = load_workbook(excel, data_only=True)
    ws = wb[wb.sheetnames[0]]
    encontrados: List[str] = []
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        nombre = _normalizar_id_experimento(header)
        if _es_id_experimento(nombre):
            encontrados.append(nombre)
    unicos = sorted(set(encontrados), key=lambda x: int(str(x).split("_")[-1]))
    return unicos


def _listar_experimentos_en_metricas_excel(archivo_metricas: str) -> List[str]:
    if not os.path.isfile(archivo_metricas):
        return []
    with pd.ExcelFile(archivo_metricas, engine="openpyxl") as xl:
        if "ranking" in xl.sheet_names:
            df = xl.parse("ranking")
        elif "metricas" in xl.sheet_names:
            df = xl.parse("metricas")
        else:
            df = xl.parse(xl.sheet_names[0])
    if "experimento" not in df.columns:
        return []
    exp = df["experimento"].apply(_extraer_id_experimento_desde_nombre)
    exp = exp[exp.apply(_es_id_experimento)]
    if exp.empty:
        return []
    unicos = sorted(set(exp.tolist()), key=lambda x: int(str(x).split("_")[-1]))
    return unicos


def _extraer_id_experimento_desde_nombre(nombre: str) -> str:
    texto = str(nombre).strip()
    m = re.match(r"^(experimento_\d+)(?:__\d+)?$", texto, flags=re.IGNORECASE)
    if m:
        return _normalizar_id_experimento(m.group(1))
    return _normalizar_id_experimento(texto)


def _seleccionar_top_experimentos_desde_metricas(
    archivo_metricas: str,
    top_n: int,
    criterio_principal: str,
    experimentos_permitidos: Optional[Sequence[str]] = None,
) -> List[str]:
    if top_n <= 0:
        return []
    if not os.path.isfile(archivo_metricas):
        raise FileNotFoundError(f"Archivo de métricas no disponible: '{archivo_metricas}'")

    with pd.ExcelFile(archivo_metricas, engine="openpyxl") as xl:
        if "ranking" in xl.sheet_names:
            df = xl.parse("ranking")
        elif "metricas" in xl.sheet_names:
            df = xl.parse("metricas")
        else:
            df = xl.parse(xl.sheet_names[0])

    if "experimento" not in df.columns:
        raise ValueError("El archivo de métricas debe contener la columna 'experimento'")

    trabajo = df.copy()
    trabajo["experimento_norm"] = trabajo["experimento"].apply(_extraer_id_experimento_desde_nombre)
    trabajo = trabajo[trabajo["experimento_norm"].apply(_es_id_experimento)]

    permitidos_set: Optional[set[str]] = None
    if experimentos_permitidos is not None:
        permitidos_set = set(_normalizar_id_experimento(x) for x in experimentos_permitidos)
        trabajo = trabajo[trabajo["experimento_norm"].isin(permitidos_set)]

    if trabajo.empty:
        raise ValueError("No se encontraron experimentos válidos en el archivo de métricas")

    for col in ["score_ranking", "MAE", "RMSE", "MAPE_pct", "Pearson", "R2"]:
        if col in trabajo.columns:
            trabajo[col] = pd.to_numeric(trabajo[col], errors="coerce")

    criterio = criterio_principal if criterio_principal in trabajo.columns else ""
    if not criterio:
        if "score_ranking" in trabajo.columns:
            criterio = "score_ranking"
        elif "MAE" in trabajo.columns:
            criterio = "MAE"
        elif "RMSE" in trabajo.columns:
            criterio = "RMSE"
        else:
            raise ValueError("No hay columnas de ranking/errores disponibles para seleccionar top")

    descendentes = {"Pearson", "R2"}
    columnas_orden: List[str] = [criterio]
    ascendente: List[bool] = [criterio not in descendentes]

    for col in ["score_ranking", "MAE", "RMSE", "MAPE_pct"]:
        if col in trabajo.columns and col not in columnas_orden:
            columnas_orden.append(col)
            ascendente.append(True)
    for col in ["Pearson", "R2"]:
        if col in trabajo.columns and col not in columnas_orden:
            columnas_orden.append(col)
            ascendente.append(False)

    columnas_orden.append("experimento_norm")
    ascendente.append(True)

    ordenado = trabajo.sort_values(columnas_orden, ascending=ascendente, na_position="last")
    top = ordenado["experimento_norm"].drop_duplicates().head(top_n).tolist()

    if permitidos_set is not None and len(top) < min(top_n, len(permitidos_set)):
        faltantes = [x for x in sorted(permitidos_set, key=lambda y: int(str(y).split("_")[-1])) if x not in top]
        top.extend(faltantes[: max(0, top_n - len(top))])

    return top


def _filtrar_excel_comparacion_a_top(excel: str, top_experimentos: Sequence[str]) -> None:
    if not os.path.isfile(excel):
        raise FileNotFoundError(f"Archivo de comparación no disponible: '{excel}'")

    wb = load_workbook(excel)
    ws = wb[wb.sheetnames[0]]
    top_set = set(_normalizar_id_experimento(x) for x in top_experimentos)

    columnas_a_borrar: List[int] = []
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        texto = str(header).strip()
        texto_norm = _normalizar_id_experimento(texto)
        if _es_id_experimento(texto_norm) and texto_norm not in top_set:
            columnas_a_borrar.append(col_idx)

    for col_idx in sorted(columnas_a_borrar, reverse=True):
        ws.delete_cols(col_idx)

    wb.save(excel)


def _filtrar_resultado_metricas_a_top(archivo_metricas: str, top_experimentos: Sequence[str]) -> None:
    if not os.path.isfile(archivo_metricas):
        return

    top_set = set(_normalizar_id_experimento(x) for x in top_experimentos)
    with pd.ExcelFile(archivo_metricas, engine="openpyxl") as xl:
        hojas = {name: xl.parse(name) for name in xl.sheet_names}

    hojas_filtradas: Dict[str, pd.DataFrame] = {}
    for nombre_hoja, df in hojas.items():
        if "experimento" in df.columns:
            trabajo = df.copy()
            trabajo["experimento_norm"] = trabajo["experimento"].apply(_extraer_id_experimento_desde_nombre)
            trabajo = trabajo[trabajo["experimento_norm"].isin(top_set)]
            trabajo = trabajo.drop(columns=["experimento_norm"])
            hojas_filtradas[nombre_hoja] = trabajo
        else:
            hojas_filtradas[nombre_hoja] = df

    with pd.ExcelWriter(archivo_metricas, engine="openpyxl") as writer:
        for nombre_hoja, df in hojas_filtradas.items():
            df.to_excel(writer, sheet_name=nombre_hoja, index=False)


def _eliminar_artefactos_no_top(
    directorio_salida: str,
    archivo_log_experimentos: str,
    top_experimentos: Sequence[str],
    conservar_log_completo: bool,
) -> Dict[str, Any]:
    top_set = set(_normalizar_id_experimento(x) for x in top_experimentos)
    errores_eliminacion: List[Dict[str, str]] = []

    patron_radar = re.compile(r"^radar_departamentos_(EXPERIMENTO_\d+)\.(csv|pkl)$", flags=re.IGNORECASE)
    for nombre in os.listdir(directorio_salida):
        match = patron_radar.match(nombre)
        if not match:
            continue
        exp = _normalizar_id_experimento(match.group(1))
        if exp in top_set:
            continue
        ruta = os.path.join(directorio_salida, nombre)
        try:
            os.remove(ruta)
        except FileNotFoundError:
            pass
        except OSError as exc:
            errores_eliminacion.append({"ruta": ruta, "error": str(exc)})

    graficos_dir = os.path.join(directorio_salida, cfg.DIRECTORIO_GRAFICOS_METRICAS)
    if os.path.isdir(graficos_dir):
        for nombre in os.listdir(graficos_dir):
            ruta = os.path.join(graficos_dir, nombre)
            if not os.path.isdir(ruta):
                continue
            exp = _extraer_id_experimento_desde_nombre(nombre)
            if exp in top_set:
                continue
            try:
                shutil.rmtree(ruta)
            except FileNotFoundError:
                pass
            except OSError as exc:
                errores_eliminacion.append({"ruta": ruta, "error": str(exc)})

    if not conservar_log_completo and os.path.isfile(archivo_log_experimentos):
        lineas_filtradas: List[str] = []
        with open(archivo_log_experimentos, "r", encoding="utf-8") as f:
            for linea in f:
                texto = linea.strip()
                if not texto:
                    continue
                try:
                    payload = json.loads(texto)
                except json.JSONDecodeError:
                    continue
                exp = _normalizar_id_experimento(payload.get("EXPERIMENTO_N", ""))
                if exp in top_set:
                    payload["EXPERIMENTO_N"] = exp
                    lineas_filtradas.append(json.dumps(payload, ensure_ascii=False))

        with open(archivo_log_experimentos, "w", encoding="utf-8") as f:
            for linea in lineas_filtradas:
                f.write(linea + "\n")

    return {
        "errores_eliminacion": errores_eliminacion,
    }


def podar_experimentos_top_n(
    archivo_comparacion_excel: str,
    archivo_metricas_excel: str,
    directorio_salida: str,
    archivo_log_experimentos: str,
    top_n: int = 10,
    criterio_ranking: str = "score_ranking",
    conservar_log_completo: bool = True,
) -> Dict[str, Any]:
    if top_n <= 0:
        raise ValueError("top_n debe ser >= 1")

    experimentos_en_excel = _listar_experimentos_en_excel(archivo_comparacion_excel)
    experimentos_en_metricas = _listar_experimentos_en_metricas_excel(archivo_metricas_excel)
    set_excel = set(experimentos_en_excel)
    set_metricas = set(experimentos_en_metricas)
    experimentos_sin_metricas = sorted([x for x in set_excel if x not in set_metricas], key=lambda x: int(x.split("_")[-1]))
    experimentos_solo_metricas = sorted([x for x in set_metricas if x not in set_excel], key=lambda x: int(x.split("_")[-1]))

    if len(experimentos_en_excel) <= top_n:
        return {
            "top_experimentos": experimentos_en_excel,
            "eliminados": [],
            "criterio_ranking": criterio_ranking,
            "top_n": top_n,
            "archivo_metricas_excel": archivo_metricas_excel,
            "archivo_comparacion_excel": archivo_comparacion_excel,
            "experimentos_sin_metricas": experimentos_sin_metricas,
            "experimentos_solo_metricas": experimentos_solo_metricas,
            "errores_eliminacion": [],
        }

    top_experimentos = _seleccionar_top_experimentos_desde_metricas(
        archivo_metricas=archivo_metricas_excel,
        top_n=top_n,
        criterio_principal=criterio_ranking,
        experimentos_permitidos=experimentos_en_excel,
    )
    top_set = set(top_experimentos)
    eliminados = sorted([x for x in experimentos_en_excel if x not in top_set], key=lambda x: int(x.split("_")[-1]))

    _filtrar_excel_comparacion_a_top(archivo_comparacion_excel, top_experimentos)
    _filtrar_resultado_metricas_a_top(archivo_metricas_excel, top_experimentos)
    resultado_eliminacion = _eliminar_artefactos_no_top(
        directorio_salida=directorio_salida,
        archivo_log_experimentos=archivo_log_experimentos,
        top_experimentos=top_experimentos,
        conservar_log_completo=conservar_log_completo,
    )

    return {
        "top_experimentos": top_experimentos,
        "eliminados": eliminados,
        "criterio_ranking": criterio_ranking,
        "top_n": top_n,
        "archivo_metricas_excel": archivo_metricas_excel,
        "archivo_comparacion_excel": archivo_comparacion_excel,
        "experimentos_sin_metricas": experimentos_sin_metricas,
        "experimentos_solo_metricas": experimentos_solo_metricas,
        "errores_eliminacion": resultado_eliminacion.get("errores_eliminacion", []),
    }


def _resolver_pesos_indicadores(
    usar_pesos_fijos: bool,
    desactivar_aleatoriedad: bool,
    pesos_fijos: Optional[Dict[str, float]],
    rng: random.Random,
) -> Dict[str, float]:
    if usar_pesos_fijos or desactivar_aleatoriedad:
        base = dict(CalculadorRadar.PESOS_INDICADORES_FIJOS)
        if isinstance(pesos_fijos, dict):
            base.update(pesos_fijos)
        return CalculadorRadar.normalizar_pesos_indicadores(base)
    base = {k: rng.random() for k in CalculadorRadar.PESOS_INDICADORES_FIJOS.keys()}
    return CalculadorRadar.normalizar_pesos_indicadores(base)


def ejecutar_experimentos_radar(
    ruta_indicadores_csv: str = os.path.join(cfg.RUTA_SALIDA_PIPELINE, "indicadores_transformers_departamento.csv"),
    ruta_radar_csv: str = os.path.join(cfg.RUTA_SALIDA_PIPELINE, "radar_departamentos.csv"),
    numero_iteraciones: int = 1,
    usar_pesos_fijos: bool = False,
    pesos_fijos: Optional[Dict[str, float]] = None,
    desactivar_aleatoriedad: bool = False,
    archivo_comparacion_excel: str = cfg.ARCHIVO_COMPARACION_EXCEL,
    archivo_log_experimentos: str = "experimentos_radar.jsonl",
    operaciones_habilitadas: Optional[List[str]] = None,
    semilla: Optional[int] = None,
    directorio_salida: str = cfg.RUTA_SALIDA_PIPELINE,
    nombre_experimento_inicial: Optional[str] = None,
    aplicar_poda_top_n: bool = False,
    top_n: int = 10,
    criterio_ranking: str = "score_ranking",
    conservar_log_completo: bool = True,
    salida_metricas: Optional[str] = None,
    archivo_metricas_excel: str = "",
) -> List[Dict[str, Any]]:
    if numero_iteraciones < 1:
        raise ValueError("numero_iteraciones debe ser >= 1")
    if not os.path.isfile(ruta_indicadores_csv):
        raise FileNotFoundError(f"Archivo no disponible: '{ruta_indicadores_csv}'")
    if not os.path.isfile(ruta_radar_csv):
        raise FileNotFoundError(f"Archivo no disponible: '{ruta_radar_csv}'")

    df_indicadores = pd.read_csv(ruta_indicadores_csv)
    df_radar_base = pd.read_csv(ruta_radar_csv)
    if 'departamento' not in df_indicadores.columns:
        raise ValueError("indicadores_transformers_departamento.csv debe contener la columna 'departamento'")
    if 'departamento' not in df_radar_base.columns:
        raise ValueError("radar_departamentos.csv debe contener la columna 'departamento'")

    operaciones = list(CalculadorRadar.OPERACIONES) if operaciones_habilitadas is None else [o.strip() for o in operaciones_habilitadas if o and o.strip()]
    if not operaciones:
        raise ValueError("Debe habilitar al menos una operación")
    invalidas = [o for o in operaciones if o not in CalculadorRadar.OPERACIONES]
    if invalidas:
        raise ValueError(f"Operaciones no válidas: {invalidas}")

    os.makedirs(directorio_salida, exist_ok=True)
    calculador = CalculadorRadar()
    rng = random.Random(semilla)

    if nombre_experimento_inicial is not None:
        indice_inicial = _parsear_indice_experimento(nombre_experimento_inicial)
        if indice_inicial is None:
            raise ValueError("nombre_experimento_inicial debe tener formato 'EXPERIMENTO_N'")
        if numero_iteraciones > 1:
            indice_final = indice_inicial + numero_iteraciones - 1
            max_existente = max(
                _max_indice_experimento_en_excel(archivo_comparacion_excel),
                _max_indice_experimento_en_log(archivo_log_experimentos),
            )
            if indice_final <= max_existente:
                raise ValueError("El rango de experimentos solicitado ya existe en Excel/log")
        indice_inicio = indice_inicial
    else:
        indice_inicio = _siguiente_indice_experimento(archivo_comparacion_excel, archivo_log_experimentos)
    resultados: List[Dict[str, Any]] = []
    salida_metricas_final = salida_metricas if salida_metricas else directorio_salida
    archivo_metricas_objetivo = archivo_metricas_excel.strip() if archivo_metricas_excel and archivo_metricas_excel.strip() else os.path.join(
        salida_metricas_final,
        f"{cfg.PREFIJO_RESULTADO_COMPARACION}.xlsx",
    )

    for i in range(numero_iteraciones):
        experimento_id = f"EXPERIMENTO_{indice_inicio + i}"
        operacion = operaciones[0] if desactivar_aleatoriedad else rng.choice(operaciones)

        if operacion == CalculadorRadar.OPERACION_BLOQUES:
            pesos_utilizados = dict(CalculadorRadar.PESOS_BLOQUES_FIJOS)
            df_radar = calculador.calcular_desde_indicadores(
                df_indicadores=df_indicadores,
                df_radar_base=df_radar_base,
                operacion=operacion,
                pesos=None,
            )
        else:
            pesos_utilizados = _resolver_pesos_indicadores(
                usar_pesos_fijos=usar_pesos_fijos,
                desactivar_aleatoriedad=desactivar_aleatoriedad,
                pesos_fijos=pesos_fijos,
                rng=rng,
            )
            df_radar = calculador.calcular_desde_indicadores(
                df_indicadores=df_indicadores,
                df_radar_base=df_radar_base,
                operacion=operacion,
                pesos=pesos_utilizados,
            )

        _validar_consistencia(df_indicadores, df_radar)

        ruta_radar_experimento = os.path.join(directorio_salida, f"radar_departamentos_{experimento_id}.csv")
        ruta_radar_experimento_pkl = os.path.join(directorio_salida, f"radar_departamentos_{experimento_id}.pkl")
        ruta_radar_actual = os.path.join(directorio_salida, "radar_departamentos.csv")
        ruta_radar_actual_pkl = os.path.join(directorio_salida, "radar_departamentos.pkl")
        df_radar.to_csv(ruta_radar_experimento, index=False)
        df_radar.to_pickle(ruta_radar_experimento_pkl)
        df_radar.to_csv(ruta_radar_actual, index=False)
        df_radar.to_pickle(ruta_radar_actual_pkl)

        _actualizar_excel_experimento(archivo_comparacion_excel, df_radar, experimento_id)

        payload_log = {
            "EXPERIMENTO_N": experimento_id,
            "OPERACION_USADA": operacion,
            "PESOS": pesos_utilizados,
        }
        _append_log_experimentos(archivo_log_experimentos, payload_log)

        resultados.append(
            {
                "EXPERIMENTO_N": experimento_id,
                "OPERACION_USADA": operacion,
                "PESOS": pesos_utilizados,
                "RUTA_RADAR_EXPERIMENTO": ruta_radar_experimento,
                "RUTA_RADAR_EXPERIMENTO_PKL": ruta_radar_experimento_pkl,
                "RUTA_RADAR_ACTUAL": ruta_radar_actual,
                "RUTA_RADAR_ACTUAL_PKL": ruta_radar_actual_pkl,
                "ARCHIVO_COMPARACION_EXCEL": archivo_comparacion_excel,
                "ARCHIVO_LOG": archivo_log_experimentos,
            }
        )

    if aplicar_poda_top_n:
        resultado_metricas = mc.calcular_metricas(archivo_comparacion_excel, salida_metricas_final)
        if resultado_metricas.get("excel_salida"):
            archivo_metricas_objetivo = resultado_metricas["excel_salida"]

        resultado_poda = podar_experimentos_top_n(
            archivo_comparacion_excel=archivo_comparacion_excel,
            archivo_metricas_excel=archivo_metricas_objetivo,
            directorio_salida=directorio_salida,
            archivo_log_experimentos=archivo_log_experimentos,
            top_n=top_n,
            criterio_ranking=criterio_ranking,
            conservar_log_completo=conservar_log_completo,
        )

        for item in resultados:
            item["PODA_TOP_N_APLICADA"] = True
            item["PODA_TOP_N_RESULTADO"] = resultado_poda
            item["ARCHIVO_METRICAS_EXCEL"] = archivo_metricas_objetivo

    return resultados


def _parsear_pesos_json(texto: str) -> Optional[Dict[str, float]]:
    if not texto or not texto.strip():
        return None
    payload = json.loads(texto)
    if not isinstance(payload, dict):
        raise ValueError("--pesos-fijos-json debe ser un JSON de tipo objeto")
    return payload


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--indicadores-csv", default=os.path.join(cfg.RUTA_SALIDA_PIPELINE, "indicadores_transformers_departamento.csv"))
    parser.add_argument("--radar-csv", default=os.path.join(cfg.RUTA_SALIDA_PIPELINE, "radar_departamentos.csv"))
    parser.add_argument("--numero-iteraciones", type=int, default=1)
    parser.add_argument("--usar-pesos-fijos", action="store_true")
    parser.add_argument("--pesos-fijos-json", default="")
    parser.add_argument("--desactivar-aleatoriedad", action="store_true")
    parser.add_argument("--archivo-comparacion-excel", default=cfg.ARCHIVO_COMPARACION_EXCEL)
    parser.add_argument("--archivo-log-experimentos", default="experimentos_radar.jsonl")
    parser.add_argument("--operaciones", default="bloques,indicadores_transformers")
    parser.add_argument("--semilla", type=int, default=None)
    parser.add_argument("--directorio-salida", default=cfg.RUTA_SALIDA_PIPELINE)
    parser.add_argument("--nombre-experimento-inicial", default="")
    parser.add_argument("--aplicar-poda-top-n", action="store_true")
    parser.add_argument("--top-n", type=int, default=10)
    parser.add_argument("--criterio-ranking", default="score_ranking")
    parser.add_argument("--conservar-log-completo", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--salida-metricas", default="")
    parser.add_argument("--archivo-metricas-excel", default="")
    args = parser.parse_args()

    pesos_fijos = _parsear_pesos_json(args.pesos_fijos_json)
    operaciones = [x.strip() for x in args.operaciones.split(",") if x.strip()]
    nombre_experimento_inicial = args.nombre_experimento_inicial.strip() or None

    salida_metricas = args.salida_metricas.strip() or None

    resultados = ejecutar_experimentos_radar(
        ruta_indicadores_csv=args.indicadores_csv,
        ruta_radar_csv=args.radar_csv,
        numero_iteraciones=args.numero_iteraciones,
        usar_pesos_fijos=args.usar_pesos_fijos,
        pesos_fijos=pesos_fijos,
        desactivar_aleatoriedad=args.desactivar_aleatoriedad,
        archivo_comparacion_excel=args.archivo_comparacion_excel,
        archivo_log_experimentos=args.archivo_log_experimentos,
        operaciones_habilitadas=operaciones,
        semilla=args.semilla,
        directorio_salida=args.directorio_salida,
        nombre_experimento_inicial=nombre_experimento_inicial,
        aplicar_poda_top_n=args.aplicar_poda_top_n,
        top_n=args.top_n,
        criterio_ranking=args.criterio_ranking,
        conservar_log_completo=args.conservar_log_completo,
        salida_metricas=salida_metricas,
        archivo_metricas_excel=args.archivo_metricas_excel,
    )
    print(pd.DataFrame(resultados).to_string(index=False))


if __name__ == "__main__":
    main()
