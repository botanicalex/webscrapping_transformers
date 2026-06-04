"""
Genera comparacion_radares_V3_P5.xlsx  y  comparacion_radares_V3_P8.xlsx
con sus respectivos archivos de metricas.

P5 - NLI Puro (25 indicadores): excluye los 5 temas zero-shot y las 6 entidades booleanas.
P8 - Sin Entidades (30 indicadores): excluye solo las 6 entidades booleanas.

Estructura de cada V3:
  Departamento | radar_oficial_promedio | radar_oficial_promedio_normalizado
  | Clasificacion_radar_oficial_promedio (DANE invertida)
  | Radar_completo_promedio_normalizado | Clasificacion_Radar_completo
  | radar_actualizado_promedio_normalizado | Clasificacion_radar_actualizado
"""

import os
import shutil
from typing import Dict, List, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import (
    accuracy_score, cohen_kappa_score, confusion_matrix,
    classification_report, precision_recall_fscore_support,
)

# =========================================================================== #
# Rutas
# =========================================================================== #
_DIR = os.path.dirname(os.path.abspath(__file__))

RUTA_V2          = os.path.join(_DIR, "comparacion_radares_V2.xlsx")
RUTA_NUEVO_ALEXA = os.path.join(_DIR, "Indicadores transformers",
                                "nuevo_alexa_indicadores_transformers_departamento.csv")
RUTA_ACTUALIZADO = os.path.join(_DIR, "Indicadores transformers",
                                "Actualizado", "actualizado_indicadores_transformers_departamento.csv")

ETIQUETAS_3 = ['Bajo', 'Medio', 'Alto']

_EQUIV: Dict[str, str] = {
    'San Andres y Providencia': 'San Andres',
    'San Andres': 'San Andres',
    'San Andres y Providencia': 'San Andres',
    'San Andrés y Providencia': 'San Andrés',
}
def _norm(s: str) -> str:
    s = str(s).strip()
    return _EQUIV.get(s, s)

INVERSION: Dict[str, str] = {'Alto': 'Bajo', 'Bajo': 'Alto', 'Medio': 'Medio'}

# =========================================================================== #
# Listas de indicadores por propuesta
# =========================================================================== #

# Las 6 entidades booleanas (excluidas en P5 y P8)
ENTIDADES_BOOLEANAS = [
    'grupos_etnicos_entidades', 'grupos_armados_entidades', 'organizaciones_entidades',
    'lideres_entidades', 'instituciones_entidades', 'actores_economicos_entidades',
]

# Los 5 temas zero-shot (excluidos solo en P5)
TEMAS_ZEROSHOT = [
    'participacion_comunitaria', 'incentivos_economicos', 'fortalecimiento_institucional',
    'impactos_ambientales', 'conflictos_socioambientales',
]

# Los 25 indicadores NLI (eventos + posturas + indicadores_nli)
IND_P5_NLI_PURO = [
    # eventos (8)
    'desplazamiento_forzado', 'reasentamiento', 'protesta_social', 'amenaza_intimidacion',
    'consulta_previa', 'audiencia_publica', 'taller_participativo', 'conflicto_territorial',
    # posturas (7)
    'rechazo_proyecto', 'deficit_derechos', 'denuncia_violacion', 'deficit_participacion_efectiva',
    'ruptura_dialogo', 'reivindicacion_territorial', 'exclusion_participacion',
    # indicadores NLI (10)
    'equidad_inclusion', 'grupos_etnicos', 'movimientos_sociales', 'grupos_poblacionales_afectados',
    'participacion_economica_local', 'transparencia_contractual', 'zonas_proteccion_alimentaria',
    'respeto_territorios', 'presencia_grupos_armados', 'desaparicion_lideres',
]

# Los 30 indicadores sin entidades booleanas
IND_P8_SIN_ENTIDADES = [
    'participacion_comunitaria', 'incentivos_economicos', 'fortalecimiento_institucional',
    'impactos_ambientales', 'conflictos_socioambientales', 'desplazamiento_forzado',
    'reasentamiento', 'protesta_social', 'amenaza_intimidacion', 'consulta_previa',
    'audiencia_publica', 'taller_participativo', 'conflicto_territorial',
    'rechazo_proyecto', 'deficit_derechos', 'denuncia_violacion', 'deficit_participacion_efectiva',
    'ruptura_dialogo', 'reivindicacion_territorial', 'exclusion_participacion', 'equidad_inclusion',
    'grupos_etnicos', 'movimientos_sociales', 'grupos_poblacionales_afectados',
    'participacion_economica_local', 'transparencia_contractual', 'zonas_proteccion_alimentaria',
    'respeto_territorios', 'presencia_grupos_armados', 'desaparicion_lideres',
]


# =========================================================================== #
# Funciones base
# =========================================================================== #
def leer_radar_oficial_v2(ruta_v2: str) -> pd.DataFrame:
    wb = load_workbook(ruta_v2, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    data: Dict[str, List] = {}
    for c_idx, h in enumerate(headers, start=1):
        if h is not None:
            data[str(h).strip()] = [
                ws.cell(row=r, column=c_idx).value for r in range(2, ws.max_row + 1)
            ]
    df = pd.DataFrame(data)
    col_dep = 'Departamento' if 'Departamento' in df.columns else 'departamento'
    cols = [col_dep, 'radar_oficial_promedio', 'radar_oficial_promedio_normalizado',
            'Clasificacion_radar_oficial_promedio']
    df_out = df[cols].copy().rename(columns={col_dep: 'Departamento'})
    df_out = df_out[df_out['Departamento'].notna()].reset_index(drop=True)
    df_out['Departamento'] = df_out['Departamento'].astype(str).str.strip().apply(_norm)
    # Invertir clasificacion oficial (Alto -> Bajo para que Alto = alto riesgo)
    df_out['Clasificacion_radar_oficial_promedio'] = (
        df_out['Clasificacion_radar_oficial_promedio'].astype(str).map(INVERSION)
    )
    return df_out


def calcular_radar(df_ind: pd.DataFrame, indicadores: List[str],
                   nombre: str) -> pd.DataFrame:
    """
    Promedio simple de los indicadores dados -> min-max 0-1 -> terciles.
    Alto valor -> 'Alto' (alto riesgo). Sin inversion de etiquetas.
    """
    df = df_ind.copy()
    df['departamento'] = df['departamento'].astype(str).str.strip().apply(_norm)
    faltantes = [c for c in indicadores if c not in df.columns]
    if faltantes:
        print(f"  [{nombre}] Indicadores faltantes (->0): {faltantes}")
        for c in faltantes:
            df[c] = 0.0
    for c in indicadores:
        df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0.0)

    raw = df[indicadores].mean(axis=1)
    mn, mx = raw.min(), raw.max()
    norm = (raw - mn) / (mx - mn) if mx > mn else pd.Series([0.5] * len(raw), index=raw.index)
    clasif = pd.qcut(norm, q=3, labels=ETIQUETAS_3, duplicates='drop').astype(str)

    out = pd.DataFrame({
        'departamento': df['departamento'].values,
        f'{nombre}_promedio_normalizado': norm.values,
        f'Clasificacion_{nombre}': clasif.values,
    })
    print(f"  [{nombre}] {len(df)} deptos | bins: "
          f"{norm.quantile([0, .33, .67, 1]).round(3).tolist()}")
    return out


def construir_v3_variante(indicadores: List[str], nombre_variante: str,
                           ruta_salida: str) -> pd.DataFrame:
    df_oficial = leer_radar_oficial_v2(RUTA_V2)

    df_alexa  = pd.read_csv(RUTA_NUEVO_ALEXA)
    df_completo = calcular_radar(df_alexa, indicadores, 'Radar_completo')
    df_completo = df_completo.rename(columns={'departamento': 'Departamento'})

    df_act_ind = pd.read_csv(RUTA_ACTUALIZADO)
    df_actualizado = calcular_radar(df_act_ind, indicadores, 'radar_actualizado')
    df_actualizado = df_actualizado.rename(columns={'departamento': 'Departamento'})

    df_v3 = (df_oficial
             .merge(df_completo,   on='Departamento', how='left')
             .merge(df_actualizado, on='Departamento', how='left'))

    df_v3.to_excel(ruta_salida, index=False, engine='openpyxl')
    print(f"  Guardado: {ruta_salida}  ({df_v3.shape})")
    return df_v3


def _metricas_par(y_true: pd.Series, y_pred: pd.Series,
                   etiqueta: str) -> Tuple[dict, pd.DataFrame, np.ndarray]:
    mask = (y_true.notna() & y_pred.notna()
            & (y_true.astype(str).str.strip() != '')
            & (y_pred.astype(str).str.strip() != '')
            & (y_true.astype(str) != 'nan')
            & (y_pred.astype(str) != 'nan'))
    yt = y_true[mask].astype(str).tolist()
    yp = y_pred[mask].astype(str).tolist()
    if not yt:
        return None, None, None
    acc = accuracy_score(yt, yp)
    p, r, f, _ = precision_recall_fscore_support(yt, yp, labels=ETIQUETAS_3,
                                                  average='macro', zero_division=0)
    kappa = cohen_kappa_score(yt, yp, labels=ETIQUETAS_3)
    resumen = {'comparacion': etiqueta, 'n_departamentos': len(yt),
               'accuracy': round(float(acc), 4), 'precision_macro': round(float(p), 4),
               'recall_macro': round(float(r), 4), 'f1_macro': round(float(f), 4),
               'cohen_kappa': round(float(kappa), 4)}
    rep = classification_report(yt, yp, labels=ETIQUETAS_3, target_names=ETIQUETAS_3,
                                 output_dict=True, zero_division=0)
    det = [{'comparacion': etiqueta, 'clase': c,
             'precision': round(float(rep[c]['precision']), 4),
             'recall': round(float(rep[c]['recall']), 4),
             'f1-score': round(float(rep[c]['f1-score']), 4),
             'support': int(rep[c]['support'])} for c in ETIQUETAS_3]
    cm = confusion_matrix(yt, yp, labels=ETIQUETAS_3)
    return resumen, pd.DataFrame(det), cm


def calcular_metricas(ruta_v3: str, ruta_salida: str) -> None:
    df = pd.read_excel(ruta_v3, engine='openpyxl')
    col_oficial  = 'Clasificacion_radar_oficial_promedio'
    col_completo = 'Clasificacion_Radar_completo'
    col_act      = 'Clasificacion_radar_actualizado'

    pares = [(f'Radar_completo_vs_radar_oficial', col_completo, col_oficial),
             (f'radar_actualizado_vs_radar_oficial', col_act,      col_oficial)]

    resumenes, detalles, matrices = [], [], {}
    for etq, cp, ct in pares:
        res, det, cm = _metricas_par(df[ct], df[cp], etq)
        if res:
            resumenes.append(res); detalles.append(det); matrices[etq] = cm

    df_res = pd.DataFrame(resumenes)
    df_det = pd.concat(detalles, ignore_index=True) if detalles else pd.DataFrame()
    filas_cm = []
    for etq, cm in matrices.items():
        filas_cm.append({'comparacion': etq, 'verdadero \\ predicho': '---',
                          **{c: '' for c in ETIQUETAS_3}})
        for i, cr in enumerate(ETIQUETAS_3):
            fila = {'comparacion': etq, 'verdadero \\ predicho': cr}
            for j, cp2 in enumerate(ETIQUETAS_3):
                fila[cp2] = int(cm[i, j])
            filas_cm.append(fila)
        filas_cm.append({'comparacion': '', 'verdadero \\ predicho': ''})

    with pd.ExcelWriter(ruta_salida, engine='openpyxl') as w:
        df_res.to_excel(w, sheet_name='resumen', index=False)
        df_det.to_excel(w, sheet_name='detalle_por_clase', index=False)
        pd.DataFrame(filas_cm).to_excel(w, sheet_name='matrices_confusion', index=False)

    print(f"  Metricas: {ruta_salida}")
    print(df_res.to_string(index=False))


# =========================================================================== #
# Main
# =========================================================================== #
def main():
    variantes = [
        {
            'nombre': 'P5_NLI_Puro',
            'indicadores': IND_P5_NLI_PURO,
            'v3': os.path.join(_DIR, "comparacion_radares_V3_P5.xlsx"),
            'metricas': os.path.join(_DIR, "metricas_clasificacion_V3_P5.xlsx"),
        },
        {
            'nombre': 'P8_Sin_Entidades',
            'indicadores': IND_P8_SIN_ENTIDADES,
            'v3': os.path.join(_DIR, "comparacion_radares_V3_P8.xlsx"),
            'metricas': os.path.join(_DIR, "metricas_clasificacion_V3_P8.xlsx"),
        },
    ]

    for v in variantes:
        print(f"\n{'='*60}")
        print(f"GENERANDO {v['nombre']} ({len(v['indicadores'])} indicadores)")
        print(f"{'='*60}")
        construir_v3_variante(v['indicadores'], v['nombre'], v['v3'])
        calcular_metricas(v['v3'], v['metricas'])

    print("\nListo. V3 original (P0) no fue modificado.")


if __name__ == '__main__':
    main()
