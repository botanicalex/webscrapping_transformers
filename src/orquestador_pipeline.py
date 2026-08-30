# -*- coding: utf-8 -*-
import argparse
import json
import os
import sys
from typing import Any, Dict, List, Optional, Tuple
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

import config_pipeline as cfg
import metricas_y_calculo_de_error as mc
import Transformer_optimo as tf
import radar as rd


# ---------------------------------------------------------------------------
# Helpers internos
# ---------------------------------------------------------------------------

def _ejecutar_etapa(nombre: str, fn: Any) -> Any:
    print(f"\n[{nombre.upper()}] Iniciando etapa...")
    try:
        resultado = fn()
        print(f"[{nombre.upper()}] Etapa completada.")
        return resultado
    except Exception as exc:
        print(f"\n[ERROR] Etapa '{nombre}' falló: {exc}", file=sys.stderr)
        sys.exit(1)


# ---------------------------------------------------------------------------
# Etapas del pipeline
# ---------------------------------------------------------------------------

def ejecutar_pipeline(skip_scraping: bool, ruta_pkl: str, salida: str) -> Tuple[str, str, str]:
    if not skip_scraping:
        tf.correr_scraping(tf.sc.FECHA_DESDE, tf.sc.FECHA_HASTA, ruta_pkl, temas=None)

    ruta_procesado_pkl, _ = tf.run_pipeline_transformers(ruta_pkl, salida)
    ruta_indicadores_csv = os.path.join(salida, "indicadores_transformers_departamento.csv")
    ruta_radar_csv = os.path.join(salida, "radar_departamentos.csv")
    resultados_radar = ejecutar_radar_experimentos(
        ruta_indicadores_csv=ruta_indicadores_csv,
        ruta_radar_csv=ruta_radar_csv,
        excel=cfg.ARCHIVO_COMPARACION_EXCEL,
        salida=salida,
        nombre_experimento="experimento_1",
        numero_iteraciones=1,
        desactivar_aleatoriedad=True,
        operaciones_radar=[rd.CalculadorRadar.OPERACION_BLOQUES],
    )
    ruta_radar_pkl = resultados_radar[-1]["RUTA_RADAR_ACTUAL_PKL"]
    return ruta_procesado_pkl, ruta_radar_pkl, ruta_radar_csv


def _normalizar_nombre_experimento(nombre_experimento: str) -> str:
    if not nombre_experimento.startswith("experimento_"):
        return f"experimento_{nombre_experimento}"
    return nombre_experimento


def _es_celda_vacia(valor: Any) -> bool:
    if valor is None:
        return True
    if isinstance(valor, float) and pd.isna(valor):
        return True
    return str(valor).strip() == ""


def _detectar_columna_vacia_en_hoja(ws: Worksheet) -> Optional[int]:
    max_col = ws.max_column
    max_row = ws.max_row
    for col_idx in range(1, max_col + 1):
        header = ws.cell(row=1, column=col_idx).value
        header_txt = "" if header is None else str(header).strip()
        es_header_vacio = header_txt == "" or header_txt.lower().startswith("unnamed:")
        if not es_header_vacio:
            continue
        if max_row <= 1:
            return col_idx
        columna_vacia = True
        for row_idx in range(2, max_row + 1):
            if not _es_celda_vacia(ws.cell(row=row_idx, column=col_idx).value):
                columna_vacia = False
                break
        if columna_vacia:
            return col_idx
    return None


def _asegurar_excel_base(excel: str, departamentos: pd.Series) -> str:
    if os.path.isfile(excel):
        wb = load_workbook(excel)
        return wb.sheetnames[0]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="departamento")
    ws.cell(row=1, column=2, value="radar_oficial_promedio")
    dep_vals = departamentos.dropna().astype(str).str.strip().drop_duplicates().tolist()
    for i, dep in enumerate(dep_vals, start=2):
        ws.cell(row=i, column=1, value=dep)
        ws.cell(row=i, column=2, value="None")
    wb.save(excel)
    return ws.title


def actualizar_excel_experimento(excel: str, radar_csv: str, nombre_experimento: str) -> None:
    nombre_experimento = _normalizar_nombre_experimento(nombre_experimento)

    if not os.path.isfile(radar_csv):
        raise FileNotFoundError(f"No existe el CSV de radar: '{radar_csv}'")

    df_radar = pd.read_csv(radar_csv)
    if "departamento" not in df_radar.columns or "radar_propio" not in df_radar.columns:
        raise ValueError(
            "El CSV de radar debe contener las columnas 'departamento' y 'radar_propio'"
        )

    df_radar = df_radar[["departamento", "radar_propio"]].copy()
    df_radar["departamento"] = df_radar["departamento"].astype(str).str.strip()
    df_radar["radar_propio"] = pd.to_numeric(df_radar["radar_propio"], errors="coerce")

    sheet_name = _asegurar_excel_base(excel, df_radar["departamento"])

    with pd.ExcelFile(excel, engine="openpyxl") as xl:
        df_excel = xl.parse(sheet_name)

    if "departamento" not in df_excel.columns:
        raise ValueError("El Excel debe contener la columna 'departamento'")

    df_excel["departamento"] = df_excel["departamento"].astype(str).str.strip()

    if nombre_experimento in df_excel.columns:
        raise ValueError(
            f"La columna '{nombre_experimento}' ya existe en '{excel}'. Usa otro nombre de experimento."
        )

    radar_map = {
        k.casefold(): v
        for k, v in zip(df_radar["departamento"], df_radar["radar_propio"])
    }
    valores_experimento = df_excel["departamento"].map(
        lambda d: radar_map.get(str(d).casefold())
    )

    wb = load_workbook(excel)
    ws = wb[sheet_name]

    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        if str(header).strip() == nombre_experimento:
            raise ValueError(
                f"La columna '{nombre_experimento}' ya existe en '{excel}'. Usa otro nombre de experimento."
            )

    col_objetivo = _detectar_columna_vacia_en_hoja(ws)
    if col_objetivo is None:
        col_objetivo = ws.max_column + 1

    ws.cell(row=1, column=col_objetivo, value=nombre_experimento)

    idx_departamento = {
        str(ws.cell(row=row_idx, column=1).value).strip().casefold(): row_idx
        for row_idx in range(2, ws.max_row + 1)
    }

    faltantes = 0
    for dep, valor in zip(df_excel["departamento"], valores_experimento):
        row_idx = idx_departamento.get(str(dep).casefold())
        if row_idx is None:
            continue
        if pd.isna(valor):
            ws.cell(row=row_idx, column=col_objetivo, value="None")
            faltantes += 1
        else:
            ws.cell(row=row_idx, column=col_objetivo, value=float(valor))

    wb.save(excel)

    print(
        f"Excel actualizado en '{excel}' con columna '{nombre_experimento}'. "
        f"Departamentos sin valor: {faltantes}"
    )


def ejecutar_radar_experimentos(
    ruta_indicadores_csv: str,
    ruta_radar_csv: str,
    excel: str,
    salida: str,
    nombre_experimento: str,
    numero_iteraciones: int = 1,
    usar_pesos_fijos: bool = False,
    pesos_fijos: Optional[Dict[str, float]] = None,
    desactivar_aleatoriedad: bool = False,
    operaciones_radar: Optional[List[str]] = None,
    semilla_radar: Optional[int] = None,
    archivo_log_experimentos: str = "",
    aplicar_poda_top_n: bool = False,
    top_n: int = 10,
    criterio_ranking: str = "score_ranking",
    conservar_log_completo: bool = True,
    salida_metricas: Optional[str] = None,
    archivo_metricas_excel: str = "",
) -> List[Dict[str, Any]]:
    nombre_normalizado = _normalizar_nombre_experimento(nombre_experimento)
    sufijo = nombre_normalizado[len("experimento_"):]
    experimento_id = f"EXPERIMENTO_{sufijo}"
    archivo_log = archivo_log_experimentos.strip() if archivo_log_experimentos else ""
    if not archivo_log:
        archivo_log = os.path.join(salida, "experimentos_radar.jsonl")
    operaciones = None
    if operaciones_radar is not None:
        operaciones = [o.strip() for o in operaciones_radar if o and o.strip()]
    resultados = rd.ejecutar_experimentos_radar(
        ruta_indicadores_csv=ruta_indicadores_csv,
        ruta_radar_csv=ruta_radar_csv,
        numero_iteraciones=numero_iteraciones,
        usar_pesos_fijos=usar_pesos_fijos,
        pesos_fijos=pesos_fijos,
        desactivar_aleatoriedad=desactivar_aleatoriedad,
        archivo_comparacion_excel=excel,
        archivo_log_experimentos=archivo_log,
        operaciones_habilitadas=operaciones,
        semilla=semilla_radar,
        directorio_salida=salida,
        nombre_experimento_inicial=experimento_id,
        aplicar_poda_top_n=aplicar_poda_top_n,
        top_n=top_n,
        criterio_ranking=criterio_ranking,
        conservar_log_completo=conservar_log_completo,
        salida_metricas=salida_metricas,
        archivo_metricas_excel=archivo_metricas_excel,
    )
    if not resultados:
        raise ValueError("No se generaron resultados de radar para el experimento")
    return resultados


def ejecutar_metricas(ruta_excel: str, salida: str) -> Dict[str, Any]:
    """LEGADO: lee desde Excel ancho de todos los experimentos."""
    return mc.calcular_metricas(ruta_excel, salida)


def ejecutar_metricas_experimento(
    experimento_id: str,
    ruta_radar_csv: str,
    ruta_v3: str,
    dir_experimento: str,
    ruta_metricas_acumulado: str,
) -> Dict[str, Any]:
    """Genera comparacion individual (V3) + actualiza metricas_experimentos.xlsx."""
    return mc.calcular_metricas_experimento(
        experimento_id=experimento_id,
        ruta_radar_csv=ruta_radar_csv,
        ruta_v3=ruta_v3,
        dir_experimento=dir_experimento,
        ruta_metricas_acumulado=ruta_metricas_acumulado,
    )


# ---------------------------------------------------------------------------
# Resumen y criterio de parada
# ---------------------------------------------------------------------------

def generar_resumen(
    ruta_radar_csv: str,
    ruta_procesado_pkl: str,
    ruta_radar_pkl: str,
    excel_comparacion: str,
    resultado_metricas: Dict[str, Any],
) -> Dict[str, Any]:
    articulos_por_departamento: Dict[str, int] = {}
    n_departamentos = 0

    if os.path.isfile(ruta_radar_csv):
        df_radar = pd.read_csv(ruta_radar_csv)
        n_departamentos = int(len(df_radar))
        if "departamento" in df_radar.columns and "n_articulos" in df_radar.columns:
            articulos_por_departamento = {
                str(dep): int(n)
                for dep, n in zip(df_radar["departamento"], df_radar["n_articulos"])
            }

    metricas_clave = [
        {
            "experimento": m.get("comparacion", m.get("experimento", "")),
            "accuracy": round(float(m["accuracy"]), 4),
            "f1_macro": round(float(m["f1_macro"]), 4),
            "cohen_kappa": round(float(m["cohen_kappa"]), 4),
        }
        for m in resultado_metricas.get("metricas", [])
    ]

    resumen: Dict[str, Any] = {
        "n_departamentos": n_departamentos,
        "articulos_por_departamento": articulos_por_departamento,
        "metricas": metricas_clave,
        "artefactos": {
            "df_procesado_pkl": ruta_procesado_pkl,
            "radar_pkl": ruta_radar_pkl,
            "radar_csv": ruta_radar_csv,
            "excel_comparacion": excel_comparacion,
            "excel_resultado": resultado_metricas.get("excel_metricas", resultado_metricas.get("excel_salida", "")),
            "graficos_dir": resultado_metricas.get("dir_experimento", resultado_metricas.get("graficos_dir", "")),
        },
    }

    print("\n========== RESUMEN FINAL ==========")
    print(json.dumps(resumen, ensure_ascii=False, indent=2))

    return resumen


def evaluar_criterio_parada(
    resultado_metricas: Dict[str, Any],
    umbral_accuracy: float = 0.70,
    umbral_mape: float = 0.0,        # ignorado — conservado por compatibilidad
    umbral_error_max: float = 0.0,   # ignorado — conservado por compatibilidad
) -> bool:
    """
    Evalúa el criterio de parada del ciclo de optimización.
    La métrica principal es accuracy de clasificación (Bajo/Medio/Alto por terciles).
    Se alcanza cuando el mejor experimento supera umbral_accuracy.
    """
    metricas = resultado_metricas.get("metricas", [])
    if not metricas:
        print("\n[CRITERIO DE PARADA] Sin datos de métricas para evaluar.")
        return False

    best_acc = max(m["accuracy"] for m in metricas)
    alcanzado = best_acc >= umbral_accuracy

    print("\n========== CRITERIO DE PARADA ==========")
    for m in metricas:
        estado = "OK" if m["accuracy"] >= umbral_accuracy else "NO OK"
        etiqueta = m.get("comparacion", m.get("experimento", ""))
        print(
            f"  {etiqueta}: Accuracy={m['accuracy']*100:.1f}%  "
            f"F1={m['f1_macro']:.3f}  kappa={m['cohen_kappa']:.3f}  "
            f"(umbral>={umbral_accuracy*100:.0f}%) [{estado}]"
        )

    if alcanzado:
        print(">>> CRITERIO DE PARADA ALCANZADO — ciclo de optimización puede detenerse.")
    else:
        print(">>> CICLO CONTINÚA — ajustar configuración y relanzar.")

    return alcanzado


# ---------------------------------------------------------------------------
# Punto de entrada
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Orquestador del pipeline radar de prensa"
    )
    parser.add_argument(
        "--only",
        choices=["scraping", "transformers", "excel", "metricas", "todo"],
        default="todo",
        help=(
            "Etapa única a ejecutar. "
            "'todo' corre el pipeline completo (default). "
            "Usar 'scraping', 'transformers', 'excel' o 'metricas' para iterar una sola etapa."
        ),
    )
    parser.add_argument(
        "--skip-scraping",
        action="store_true",
        help="En modo 'todo', omitir la etapa de scraping",
    )
    parser.add_argument("--ruta-pkl", default=cfg.RUTA_CORPUS_PKL)
    parser.add_argument("--salida", default=cfg.RUTA_SALIDA_PIPELINE)
    parser.add_argument("--excel", default=cfg.ARCHIVO_COMPARACION_EXCEL)
    parser.add_argument("--nombre-experimento", default="experimento_1")
    parser.add_argument("--iteraciones-radar", type=int, default=1)
    parser.add_argument("--usar-pesos-fijos", action="store_true")
    parser.add_argument("--pesos-fijos-json", default="")
    parser.add_argument("--desactivar-aleatoriedad", action="store_true")
    parser.add_argument("--operaciones-radar", default="bloques,indicadores_transformers")
    parser.add_argument("--semilla-radar", type=int, default=None)
    parser.add_argument("--log-experimentos-radar", default="")
    parser.add_argument("--salida-metricas", default=".")
    parser.add_argument(
        "--umbral-accuracy",
        type=float,
        default=0.70,
        help="Umbral de accuracy (clasificación) para criterio de parada (default: 0.70)",
    )
    parser.add_argument("--aplicar-poda-top-n", action="store_true")
    parser.add_argument("--top-n", type=int, default=10)
    parser.add_argument("--criterio-ranking", default="score_ranking")
    parser.add_argument("--conservar-log-completo", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--archivo-metricas-excel", default="")
    args = parser.parse_args()

    only = args.only
    ruta_radar_csv = os.path.join(args.salida, "radar_departamentos.csv")
    ruta_indicadores_csv = os.path.join(args.salida, "indicadores_transformers_departamento.csv")
    ruta_procesado_pkl = os.path.join(args.salida, "df_procesado.pkl")
    ruta_radar_pkl = os.path.join(args.salida, "radar_departamentos.pkl")
    pesos_fijos = rd._parsear_pesos_json(args.pesos_fijos_json)
    operaciones_radar = [x.strip() for x in args.operaciones_radar.split(",") if x.strip()]

    correr_scraping = only == "scraping" or (only == "todo" and not args.skip_scraping)
    correr_transformers = only in ("transformers", "todo")
    correr_excel = only in ("excel", "todo")
    correr_metricas = only in ("metricas", "todo")

    # --- Etapa A: scraping ---
    if correr_scraping:
        _ejecutar_etapa(
            "scraping",
            lambda: tf.correr_scraping(
                tf.sc.FECHA_DESDE, tf.sc.FECHA_HASTA, args.ruta_pkl, temas=None
            ),
        )

    # --- Etapa B: transformers + radar ---
    if correr_transformers:
        resultado_pipeline: Optional[Tuple[str, str]] = _ejecutar_etapa(
            "transformers",
            lambda: tf.run_pipeline_transformers(args.ruta_pkl, args.salida),
        )
        if resultado_pipeline is not None:
            ruta_procesado_pkl, ruta_radar_pkl = resultado_pipeline

    if correr_excel:
        resultados_radar: List[Dict[str, Any]] = _ejecutar_etapa(
            "excel",
            lambda: ejecutar_radar_experimentos(
                ruta_indicadores_csv=ruta_indicadores_csv,
                ruta_radar_csv=ruta_radar_csv,
                excel=args.excel,
                salida=args.salida,
                nombre_experimento=args.nombre_experimento,
                numero_iteraciones=args.iteraciones_radar,
                usar_pesos_fijos=args.usar_pesos_fijos,
                pesos_fijos=pesos_fijos,
                desactivar_aleatoriedad=args.desactivar_aleatoriedad,
                operaciones_radar=operaciones_radar,
                semilla_radar=args.semilla_radar,
                archivo_log_experimentos=args.log_experimentos_radar,
                aplicar_poda_top_n=args.aplicar_poda_top_n,
                top_n=args.top_n,
                criterio_ranking=args.criterio_ranking,
                conservar_log_completo=args.conservar_log_completo,
                salida_metricas=args.salida_metricas,
                archivo_metricas_excel=args.archivo_metricas_excel,
            ),
        )
        if resultados_radar:
            ruta_radar_csv = resultados_radar[-1]["RUTA_RADAR_ACTUAL"]
            ruta_radar_pkl = resultados_radar[-1]["RUTA_RADAR_ACTUAL_PKL"]

    # --- Etapa D: métricas + resumen + criterio de parada ---
    if correr_metricas:
        resultado_metricas: Optional[Dict[str, Any]] = _ejecutar_etapa(
            "metricas",
            lambda: ejecutar_metricas(args.excel, args.salida_metricas),
        )
        if resultado_metricas is not None:
            generar_resumen(
                ruta_radar_csv=ruta_radar_csv,
                ruta_procesado_pkl=ruta_procesado_pkl,
                ruta_radar_pkl=ruta_radar_pkl,
                excel_comparacion=args.excel,
                resultado_metricas=resultado_metricas,
            )
            evaluar_criterio_parada(
                resultado_metricas, umbral_accuracy=args.umbral_accuracy
            )


if __name__ == "__main__":
    main()
