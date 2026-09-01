import argparse
import os
import re
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler

import config_pipeline as cfg


class CalculadorRadar:
    MAPEO_PERIODICO_DEPARTAMENTO = {
        'El Colombiano': 'Antioquia', 'El Diario': 'Risaralda', 'BC Noticias': 'Caldas', 'El Quindiano': 'Quindío',
        'El País Cali': 'Valle del Cauca', 'Diario Occidente': 'Valle del Cauca', 'Diario del Sur': 'Nariño',
        'Diario del Cauca': 'Cauca', 'Chocó 7 Días': 'Chocó', 'Llano al Mundo': 'Meta', 'Diario de Casanare': 'Casanare',
        'La Voz del Cinaruco': 'Arauca', 'El Morichal': 'Vichada', 'Mi Putumayo': 'Putumayo', 'El Tiempo': 'Cundinamarca',
        'La República': 'Cundinamarca', 'Portafolio': 'Cundinamarca', 'Publimetro': 'Cundinamarca', 'Las2Orillas': 'Cundinamarca',
        'El Heraldo': 'Atlántico', 'El Universal': 'Bolívar', 'El Pilón': 'Cesar', 'El Meridiano': 'Córdoba',
        'Vanguardia': 'Santander', 'Trochando Sin Fronteras': 'Arauca', 'Enlace Television': 'Santander', 'Corrillos': 'Santander'
    }

    BLOQUES_PCA = {
        'bloque_A': ['violacion_derechos_humanos', 'irregularidad_contractual', 'conflicto_territorial'],
        'bloque_B': ['presencia_grupos_armados', 'amenaza_lideres', 'amenaza_intimidacion', 'conflictos_socioambientales'],
        'bloque_C': ['debilidad_institucional', 'conflicto_activo', 'incentivos_economicos_inequitativos'],
        'bloque_D': ['desplazamiento_forzado', 'reasentamiento', 'poblacion_afectada',
                     'zonas_proteccion_alimentaria', 'dano_territorios', 'exclusion_servicios_derechos', 'grupos_etnicos_existentes',
                     'danos_ambientales', 'derechos_vulnerados', 'resistencia_territorial'],
        'bloque_E': ['deficit_participacion_comunitaria',
                     'exclusion_comunidades', 'movimientos_sociales', 'rechazo_proyecto', 'protesta_social', 'exclusion_beneficios_economicos'],
    }

    VARS_INVERTIR: set = set()  # todos los indicadores tienen hipótesis de déficit/riesgo

    COLUMNAS_BINARIAS = [
        'deficit_participacion_comunitaria', 'incentivos_economicos_inequitativos', 'debilidad_institucional', 'danos_ambientales', 'conflictos_socioambientales',
        'desplazamiento_forzado', 'reasentamiento', 'protesta_social', 'amenaza_intimidacion', 'conflicto_territorial',
        'rechazo_proyecto', 'derechos_vulnerados', 'violacion_derechos_humanos', 'conflicto_activo', 'resistencia_territorial', 'exclusion_comunidades',
        'exclusion_servicios_derechos', 'movimientos_sociales', 'poblacion_afectada', 'exclusion_beneficios_economicos', 'irregularidad_contractual',
        'zonas_proteccion_alimentaria', 'dano_territorios', 'presencia_grupos_armados', 'amenaza_lideres',
        'grupos_etnicos_existentes',
    ]

    # Cortes fijos Bajo/Medio/Alto del radar V2 -- ver config_pipeline.py
    # (fuente unica, la comparten radar.py y metricas_y_calculo_de_error.py
    # sin crear un import circular entre los dos).
    CORTE_BAJO_MEDIO = cfg.CORTE_BAJO_MEDIO_RADAR
    CORTE_MEDIO_ALTO = cfg.CORTE_MEDIO_ALTO_RADAR

    def calcular(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        if 'departamento' not in df.columns or not df['departamento'].notna().any():
            df['departamento'] = df['periodico'].map(self.MAPEO_PERIODICO_DEPARTAMENTO)
            df = df.dropna(subset=['departamento'])
        cols = [c for c in self.COLUMNAS_BINARIAS if c in df.columns]
        for c in cols:
            if df[c].dtype == bool:
                df[c] = df[c].astype(float)
            else:
                df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0.0)
        df_tasas = df.groupby('departamento')[cols].agg('max')
        n_articulos_serie = df.groupby('departamento').size()
        for c in [c for c in cols if c in self.VARS_INVERTIR]:
            df_tasas[c] = 1 - df_tasas[c]

        radar_raw = df_tasas[cols].mean(axis=1)

        df_sub = pd.DataFrame(index=df_tasas.index)
        df_sub['n_articulos'] = n_articulos_serie.reindex(df_sub.index)
        for b, variables in self.BLOQUES_PCA.items():
            v = [x for x in variables if x in df_tasas.columns]
            df_sub[b] = df_tasas[v].mean(axis=1) * 100.0 if v else 0.0

        df_sub['corrupcion_score'] = df_sub[['bloque_A', 'bloque_B', 'bloque_C']].mean(axis=1)
        df_sub['vulneracion_score'] = df_sub[['bloque_D', 'bloque_E']].mean(axis=1)
        df_sub['radar_propio'] = radar_raw.round(4)
        df_sub['categoria_riesgo'] = self._categoria_cortes_fijos(df_sub['radar_propio'])
        out = df_sub.reset_index().rename(columns={'index': 'departamento'})
        return out[['departamento', 'n_articulos', 'bloque_A', 'bloque_B', 'bloque_C', 'bloque_D', 'bloque_E', 'corrupcion_score', 'vulneracion_score', 'radar_propio', 'categoria_riesgo']].sort_values('radar_propio', ascending=False)

    def calcular_desde_indicadores(
        self,
        df_indicadores: pd.DataFrame,
        df_radar_base: Optional[pd.DataFrame] = None,
    ) -> pd.DataFrame:
        df_tasas = self._preparar_tasas_indicadores(df_indicadores)
        n_articulos = self._resolver_n_articulos(df_radar_base, df_tasas.index)
        return self._calcular_bloques_desde_tasas(df_tasas, n_articulos)

    def _preparar_tasas_indicadores(self, df_indicadores: pd.DataFrame) -> pd.DataFrame:
        if 'departamento' not in df_indicadores.columns:
            raise ValueError("El archivo de indicadores debe contener la columna 'departamento'")
        df = df_indicadores.copy()
        df['departamento'] = df['departamento'].astype(str).str.strip()
        df = df[df['departamento'] != ""]
        for c in self.COLUMNAS_BINARIAS:
            if c not in df.columns:
                df[c] = np.nan
            df[c] = pd.to_numeric(df[c], errors='coerce')
        return df.groupby('departamento')[self.COLUMNAS_BINARIAS].mean().fillna(0.0)

    def _resolver_n_articulos(self, df_radar_base: Optional[pd.DataFrame], index_departamentos: pd.Index) -> pd.Series:
        n_articulos = pd.Series([np.nan] * len(index_departamentos), index=index_departamentos, dtype=float)
        if df_radar_base is None or df_radar_base.empty:
            return n_articulos
        if 'departamento' not in df_radar_base.columns or 'n_articulos' not in df_radar_base.columns:
            return n_articulos
        base = df_radar_base[['departamento', 'n_articulos']].copy()
        base['departamento'] = base['departamento'].astype(str).str.strip()
        base['n_articulos'] = pd.to_numeric(base['n_articulos'], errors='coerce')
        mapa = base.dropna(subset=['departamento']).drop_duplicates(subset=['departamento']).set_index('departamento')['n_articulos']
        valores = index_departamentos.to_series().map(mapa)
        return valores.astype(float)

    def _calcular_bloques_desde_tasas(self, df_tasas: pd.DataFrame, n_articulos: pd.Series) -> pd.DataFrame:
        """
        Camino de producción por defecto (V2, promovido 2026-08-31; esta rama
        usa MAX en vez de P75). `df_tasas` viene de
        `indicadores_transformers_departamento.csv`
        (`exportar_indicadores_transformers_por_departamento`), que ya trae el
        MAX por indicador y departamento — una fila por departamento, así que
        el groupby/mean de `_preparar_tasas_indicadores` es un no-op. Aquí solo
        se promedia entre los 26 indicadores (sin pesos) y se clasifica con
        cortes fijos, sin calibración z-score (ver `CORTE_BAJO_MEDIO`/
        `CORTE_MEDIO_ALTO` — están calibrados sobre la escala MAX de esta
        rama, no sobre la escala z-score, que además es monótona y no
        cambiaba la clasificación).
        """
        df_tasas = df_tasas.copy()
        cols = [c for c in self.COLUMNAS_BINARIAS if c in df_tasas.columns]
        for c in [c for c in cols if c in self.VARS_INVERTIR]:
            df_tasas[c] = 1 - df_tasas[c]

        radar_raw = df_tasas[cols].mean(axis=1)

        df_sub = pd.DataFrame(index=df_tasas.index)
        df_sub['n_articulos'] = n_articulos.reindex(df_sub.index)
        for b, variables in self.BLOQUES_PCA.items():
            v = [x for x in variables if x in df_tasas.columns]
            df_sub[b] = df_tasas[v].mean(axis=1) * 100.0 if v else 0.0

        df_sub['corrupcion_score'] = df_sub[['bloque_A', 'bloque_B', 'bloque_C']].mean(axis=1)
        df_sub['vulneracion_score'] = df_sub[['bloque_D', 'bloque_E']].mean(axis=1)
        df_sub['radar_propio'] = radar_raw.round(4)
        df_sub['categoria_riesgo'] = self._categoria_cortes_fijos(df_sub['radar_propio'])
        out = df_sub.reset_index().rename(columns={'index': 'departamento'})
        return out[['departamento', 'n_articulos', 'bloque_A', 'bloque_B', 'bloque_C', 'bloque_D', 'bloque_E', 'corrupcion_score', 'vulneracion_score', 'radar_propio', 'categoria_riesgo']].sort_values('radar_propio', ascending=False)

    @classmethod
    def _categoria_cortes_fijos(cls, s: pd.Series) -> pd.Series:
        """
        Cortes fijos Bajo/Medio/Alto (V2, ver `CORTE_BAJO_MEDIO`/
        `CORTE_MEDIO_ALTO`) — no terciles recalculados por lote: el radar debe
        poder clasificar un lugar solo (una vereda), sin otros 31 lugares con
        qué hacer terciles (contexto/08_log_decisiones.md [2026-06]).
        """
        def _clasificar(v: float) -> str:
            if pd.isna(v):
                return "None"
            if v < cls.CORTE_BAJO_MEDIO:
                return "Bajo"
            if v < cls.CORTE_MEDIO_ALTO:
                return "Medio"
            return "Alto"
        return s.apply(_clasificar)


def _es_celda_vacia(valor: Any) -> bool:
    if valor is None:
        return True
    if isinstance(valor, float) and pd.isna(valor):
        return True
    return str(valor).strip() == ""


def _detectar_columna_vacia_en_hoja(ws: Worksheet) -> Optional[int]:
    max_col = ws.max_column
    max_row = ws.max_row
    for col_idx in range(1, max_col + 1):
        header = ws.cell(row=1, column=col_idx).value
        header_txt = "" if header is None else str(header).strip()
        if header_txt != "" and not header_txt.lower().startswith("unnamed:"):
            continue
        if max_row <= 1:
            return col_idx
        columna_vacia = True
        for row_idx in range(2, max_row + 1):
            if not _es_celda_vacia(ws.cell(row=row_idx, column=col_idx).value):
                columna_vacia = False
                break
        if columna_vacia:
            return col_idx
    return None


def _asegurar_excel_base(excel: str, departamentos: pd.Series) -> str:
    if os.path.isfile(excel):
        wb = load_workbook(excel)
        return wb.sheetnames[0]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="departamento")
    ws.cell(row=1, column=2, value="radar_oficial_promedio")
    dep_vals = departamentos.dropna().astype(str).str.strip().drop_duplicates().tolist()
    for i, dep in enumerate(dep_vals, start=2):
        ws.cell(row=i, column=1, value=dep)
        ws.cell(row=i, column=2, value="None")
    wb.save(excel)
    return ws.title


def _max_indice_experimento_en_excel(excel: str) -> int:
    if not os.path.isfile(excel):
        return 0
    wb = load_workbook(excel, data_only=True)
    ws = wb[wb.sheetnames[0]]
    patron = re.compile(r"^experimento_(\d+)$", flags=re.IGNORECASE)
    max_idx = 0
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        texto = str(header).strip()
        m = patron.match(texto)
        if m:
            max_idx = max(max_idx, int(m.group(1)))
    return max_idx


def _siguiente_indice_experimento(excel: str) -> int:
    return _max_indice_experimento_en_excel(excel) + 1


def _parsear_indice_experimento(nombre_experimento: str) -> Optional[int]:
    patron = re.compile(r"^experimento_(\d+)$", flags=re.IGNORECASE)
    match = patron.match(str(nombre_experimento).strip())
    if not match:
        return None
    return int(match.group(1))


def _actualizar_excel_experimento(excel: str, df_radar: pd.DataFrame, nombre_experimento: str) -> None:
    if 'departamento' not in df_radar.columns or 'radar_propio' not in df_radar.columns:
        raise ValueError("El resultado de radar debe contener 'departamento' y 'radar_propio'")

    df_radar_local = df_radar[['departamento', 'radar_propio']].copy()
    df_radar_local['departamento'] = df_radar_local['departamento'].astype(str).str.strip()
    df_radar_local['radar_propio'] = pd.to_numeric(df_radar_local['radar_propio'], errors='coerce')

    sheet_name = _asegurar_excel_base(excel, df_radar_local['departamento'])
    with pd.ExcelFile(excel, engine="openpyxl") as xl:
        df_excel = xl.parse(sheet_name)

    if 'departamento' not in df_excel.columns:
        raise ValueError("El Excel debe contener la columna 'departamento'")

    df_excel['departamento'] = df_excel['departamento'].astype(str).str.strip()
    if nombre_experimento in df_excel.columns:
        raise ValueError(f"La columna '{nombre_experimento}' ya existe en '{excel}'")

    wb = load_workbook(excel)
    ws = wb[sheet_name]
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        if str(header).strip() == nombre_experimento:
            raise ValueError(f"La columna '{nombre_experimento}' ya existe en '{excel}'")

    idx_departamento = {
        str(ws.cell(row=row_idx, column=1).value).strip().casefold(): row_idx
        for row_idx in range(2, ws.max_row + 1)
    }
    for dep in df_radar_local['departamento'].dropna().astype(str).str.strip().tolist():
        clave = dep.casefold()
        if clave not in idx_departamento:
            nueva_fila = ws.max_row + 1
            ws.cell(row=nueva_fila, column=1, value=dep)
            if ws.max_column >= 2 and _es_celda_vacia(ws.cell(row=nueva_fila, column=2).value):
                ws.cell(row=nueva_fila, column=2, value="None")
            idx_departamento[clave] = nueva_fila

    col_objetivo = _detectar_columna_vacia_en_hoja(ws)
    if col_objetivo is None:
        col_objetivo = ws.max_column + 1
    ws.cell(row=1, column=col_objetivo, value=nombre_experimento)

    radar_map = {
        k.casefold(): v
        for k, v in zip(df_radar_local['departamento'], df_radar_local['radar_propio'])
    }
    for row_idx in range(2, ws.max_row + 1):
        dep = ws.cell(row=row_idx, column=1).value
        clave = str(dep).strip().casefold()
        valor = radar_map.get(clave)
        if pd.isna(valor):
            ws.cell(row=row_idx, column=col_objetivo, value="None")
        else:
            ws.cell(row=row_idx, column=col_objetivo, value=float(valor))

    wb.save(excel)


def _validar_consistencia(df_indicadores: pd.DataFrame, df_resultado: pd.DataFrame) -> None:
    if df_resultado.empty:
        raise ValueError("El resultado del radar está vacío")
    if 'departamento' not in df_indicadores.columns or 'departamento' not in df_resultado.columns:
        raise ValueError("Los datos de entrada y salida deben contener la columna 'departamento'")
    if 'radar_propio' not in df_resultado.columns:
        raise ValueError("La salida de radar debe contener la columna 'radar_propio'")
    deps_in = set(df_indicadores['departamento'].dropna().astype(str).str.strip().tolist())
    deps_out = set(df_resultado['departamento'].dropna().astype(str).str.strip().tolist())
    faltantes = sorted(list(deps_in - deps_out))
    if faltantes:
        raise ValueError(f"Faltan departamentos en la salida de radar: {faltantes[:10]}")
    radar_numerico = pd.to_numeric(df_resultado['radar_propio'], errors='coerce')
    if radar_numerico.isna().all():
        raise ValueError("La salida contiene radar_propio no numérico para todos los departamentos")


def ejecutar_radar_bloques(
    ruta_indicadores_csv: str = os.path.join(cfg.RUTA_SALIDA_PIPELINE, "indicadores_transformers_departamento.csv"),
    ruta_radar_csv: str = os.path.join(cfg.RUTA_SALIDA_PIPELINE, "radar_departamentos.csv"),
    numero_iteraciones: int = 1,
    archivo_comparacion_excel: str = cfg.ARCHIVO_COMPARACION_EXCEL,
    directorio_salida: str = cfg.RUTA_SALIDA_PIPELINE,
    nombre_experimento_inicial: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """
    Corre el radar (camino bloques: MAX por indicador + promedio de 26,
    cortes fijos) y registra cada corrida como una columna EXPERIMENTO_N en
    `archivo_comparacion_excel`, para que `metricas_y_calculo_de_error.py`
    pueda compararla contra el radar oficial DANE.
    """
    if numero_iteraciones < 1:
        raise ValueError("numero_iteraciones debe ser >= 1")
    if not os.path.isfile(ruta_indicadores_csv):
        raise FileNotFoundError(f"Archivo no disponible: '{ruta_indicadores_csv}'")
    if not os.path.isfile(ruta_radar_csv):
        raise FileNotFoundError(f"Archivo no disponible: '{ruta_radar_csv}'")

    df_indicadores = pd.read_csv(ruta_indicadores_csv)
    df_radar_base = pd.read_csv(ruta_radar_csv)
    if 'departamento' not in df_indicadores.columns:
        raise ValueError("indicadores_transformers_departamento.csv debe contener la columna 'departamento'")
    if 'departamento' not in df_radar_base.columns:
        raise ValueError("radar_departamentos.csv debe contener la columna 'departamento'")

    os.makedirs(directorio_salida, exist_ok=True)
    calculador = CalculadorRadar()

    if nombre_experimento_inicial is not None:
        indice_inicial = _parsear_indice_experimento(nombre_experimento_inicial)
        if indice_inicial is None:
            raise ValueError("nombre_experimento_inicial debe tener formato 'EXPERIMENTO_N'")
        if numero_iteraciones > 1:
            indice_final = indice_inicial + numero_iteraciones - 1
            max_existente = _max_indice_experimento_en_excel(archivo_comparacion_excel)
            if indice_final <= max_existente:
                raise ValueError("El rango de experimentos solicitado ya existe en Excel")
        indice_inicio = indice_inicial
    else:
        indice_inicio = _siguiente_indice_experimento(archivo_comparacion_excel)
    resultados: List[Dict[str, Any]] = []

    for i in range(numero_iteraciones):
        experimento_id = f"EXPERIMENTO_{indice_inicio + i}"
        df_radar = calculador.calcular_desde_indicadores(
            df_indicadores=df_indicadores,
            df_radar_base=df_radar_base,
        )
        _validar_consistencia(df_indicadores, df_radar)

        ruta_radar_experimento = os.path.join(directorio_salida, f"radar_departamentos_{experimento_id}.csv")
        ruta_radar_experimento_pkl = os.path.join(directorio_salida, f"radar_departamentos_{experimento_id}.pkl")
        ruta_radar_actual = os.path.join(directorio_salida, "radar_departamentos.csv")
        ruta_radar_actual_pkl = os.path.join(directorio_salida, "radar_departamentos.pkl")
        df_radar.to_csv(ruta_radar_experimento, index=False)
        df_radar.to_pickle(ruta_radar_experimento_pkl)
        df_radar.to_csv(ruta_radar_actual, index=False)
        df_radar.to_pickle(ruta_radar_actual_pkl)

        _actualizar_excel_experimento(archivo_comparacion_excel, df_radar, experimento_id)

        resultados.append(
            {
                "EXPERIMENTO_N": experimento_id,
                "RUTA_RADAR_EXPERIMENTO": ruta_radar_experimento,
                "RUTA_RADAR_EXPERIMENTO_PKL": ruta_radar_experimento_pkl,
                "RUTA_RADAR_ACTUAL": ruta_radar_actual,
                "RUTA_RADAR_ACTUAL_PKL": ruta_radar_actual_pkl,
                "ARCHIVO_COMPARACION_EXCEL": archivo_comparacion_excel,
            }
        )

    return resultados


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--indicadores-csv", default=os.path.join(cfg.RUTA_SALIDA_PIPELINE, "indicadores_transformers_departamento.csv"))
    parser.add_argument("--radar-csv", default=os.path.join(cfg.RUTA_SALIDA_PIPELINE, "radar_departamentos.csv"))
    parser.add_argument("--numero-iteraciones", type=int, default=1)
    parser.add_argument("--archivo-comparacion-excel", default=cfg.ARCHIVO_COMPARACION_EXCEL)
    parser.add_argument("--directorio-salida", default=cfg.RUTA_SALIDA_PIPELINE)
    parser.add_argument("--nombre-experimento-inicial", default="")
    args = parser.parse_args()

    nombre_experimento_inicial = args.nombre_experimento_inicial.strip() or None

    resultados = ejecutar_radar_bloques(
        ruta_indicadores_csv=args.indicadores_csv,
        ruta_radar_csv=args.radar_csv,
        numero_iteraciones=args.numero_iteraciones,
        archivo_comparacion_excel=args.archivo_comparacion_excel,
        directorio_salida=args.directorio_salida,
        nombre_experimento_inicial=nombre_experimento_inicial,
    )
    print(pd.DataFrame(resultados).to_string(index=False))


if __name__ == "__main__":
    main()
