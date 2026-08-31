# -*- coding: utf-8 -*-
"""
Métricas de clasificación para comparar experimentos vs radar oficial DANE.

Lógica de dirección (ambas fuentes):
  - radar_oficial_promedio : alto valor = ALTO riesgo  → terciles directos
  - EXPERIMENTO_N          : alto valor = ALTO riesgo  → terciles directos

Nueva arquitectura (por experimento):
  - calcular_metricas_experimento(): genera comparacion_EXP.xlsx + actualiza metricas_experimentos.xlsx
"""
import os
from typing import Any, Dict, List, Optional, Tuple

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from openpyxl import load_workbook
from sklearn.metrics import (
    accuracy_score,
    classification_report,
    cohen_kappa_score,
    confusion_matrix,
    precision_recall_fscore_support,
)

import config_pipeline as cfg

LABELS_CAT = ["Bajo", "Medio", "Alto"]
COL_CLASIFICACION_OFICIAL = "Clasificacion_radar_oficial_promedio"


# ---------------------------------------------------------------------------
# Helpers genéricos
# ---------------------------------------------------------------------------

def _to_numeric_series(serie: pd.Series) -> pd.Series:
    tokens_nulos = {"none", "null", "nan", ""}
    limpia = serie.copy()
    mascara = limpia.astype(str).str.strip().str.casefold().isin(tokens_nulos)
    limpia = limpia.mask(mascara, np.nan)
    limpia = limpia.astype(str).str.replace(",", ".", regex=False).str.strip()
    return pd.to_numeric(limpia, errors="coerce")


def _nombre_carpeta_seguro(nombre: str) -> str:
    base = str(nombre).strip() or "experimento"
    limpio = "".join(c if c.isalnum() or c in (" ", "-", "_") else "_" for c in base).strip()
    return (limpio or "experimento")[:120]


def _reemplazar_inf_para_excel(df: pd.DataFrame) -> pd.DataFrame:
    return df.replace([np.inf, -np.inf], np.nan)


# ---------------------------------------------------------------------------
# Clasificación en terciles
# ---------------------------------------------------------------------------

def _terciles(serie: pd.Series) -> pd.Series:
    """Terciles Bajo/Medio/Alto: menor valor → 'Bajo', mayor → 'Alto'."""
    n = len(serie)
    if n < 3:
        return pd.Series(["Medio"] * n, index=serie.index)
    try:
        return pd.qcut(serie.rank(method="first"), q=3, labels=LABELS_CAT).astype(str)
    except Exception:
        p33, p67 = serie.quantile(1 / 3), serie.quantile(2 / 3)
        return serie.apply(lambda x: "Bajo" if x <= p33 else ("Alto" if x >= p67 else "Medio"))


def _clasificar(serie: pd.Series) -> pd.Series:
    """Terciles: mayor valor → 'Alto' riesgo."""
    return _terciles(serie)


# ---------------------------------------------------------------------------
# Líneas base — ver contexto/07_backlog.md punto 0b y 01_objetivo_y_radar.md
# ---------------------------------------------------------------------------

def _lineas_base(y_true: pd.Series, n_articulos: Optional[pd.Series]) -> Dict[str, Any]:
    """
    Referencias contra las que una accuracy sola no dice nada:
      - azar: 1/3 analítico (3 clases equiprobables), no depende de los datos.
      - clase_mayoritaria: acertar siempre la clase más frecuente de y_true.
      - modelo_nulo_articulos: clasificar por terciles del número de artículos
        del corpus, sin leer una sola noticia. Usa el mismo departamento y el
        mismo n que la comparación real, así que es directamente comparable
        con `resumen['accuracy']` de la misma corrida.
    Cada línea se recalcula sobre la intersección vigente, no se cita de memoria.
    """
    yt = y_true.dropna()
    yt = yt[~yt.astype(str).str.strip().isin({"", "nan"})]
    n = len(yt)

    base: Dict[str, Any] = {
        "baseline_azar": round(1 / 3, 4),
        "baseline_clase_mayoritaria": None,
        "baseline_clase_mayoritaria_cual": None,
        "baseline_modelo_nulo_articulos": None,
    }
    if n == 0:
        return base

    conteo = yt.astype(str).value_counts()
    base["baseline_clase_mayoritaria"] = round(float(conteo.iloc[0] / n), 4)
    base["baseline_clase_mayoritaria_cual"] = str(conteo.index[0])

    if n_articulos is not None:
        na = pd.to_numeric(n_articulos, errors="coerce")
        comun = yt.index.intersection(na.dropna().index)
        if len(comun) >= 3:
            cat_articulos = _clasificar(na.loc[comun])
            acc_nulo = (cat_articulos.values == yt.loc[comun].astype(str).values).mean()
            base["baseline_modelo_nulo_articulos"] = round(float(acc_nulo), 4)

    return base


def _imprimir_lineas_base(base: Dict[str, Any], accuracy_actual: float) -> None:
    ancho = 42
    print("  Lineas base (no se citan de memoria, se recalculan cada corrida):")
    print(f"    {'Azar (3 clases)':<{ancho}}: {base['baseline_azar']*100:.1f}%")
    if base["baseline_clase_mayoritaria"] is not None:
        etiqueta = f"Predecir siempre '{base['baseline_clase_mayoritaria_cual']}'"
        print(f"    {etiqueta:<{ancho}}: {base['baseline_clase_mayoritaria']*100:.1f}%")
    if base["baseline_modelo_nulo_articulos"] is not None:
        print(f"    {'Modelo nulo (terciles por n_articulos)':<{ancho}}: "
              f"{base['baseline_modelo_nulo_articulos']*100:.1f}%")
    else:
        print(f"    {'Modelo nulo (terciles por n_articulos)':<{ancho}}: no disponible (falta n_articulos)")
    print(f"    {'Comparacion actual':<{ancho}}: {accuracy_actual*100:.1f}%")


def _normalizar_minmax(s: pd.Series) -> pd.Series:
    """Min-max normalization [0, 1] sobre valores no-NaN."""
    mn, mx = s.min(), s.max()
    if pd.isna(mn) or pd.isna(mx) or mx == mn:
        return pd.Series(0.5, index=s.index, dtype=float)
    return ((s - mn) / (mx - mn)).round(6)


# ---------------------------------------------------------------------------
# Métricas de clasificación para un par
# ---------------------------------------------------------------------------

def _metricas_par(
    y_true: pd.Series,
    y_pred: pd.Series,
    etiqueta: str,
) -> Tuple[Optional[dict], Optional[pd.DataFrame], Optional[np.ndarray]]:
    """Accuracy, Precision/Recall/F1 por clase, Kappa y matriz de confusión."""
    mask = (
        y_true.notna() & y_pred.notna()
        & ~y_true.astype(str).str.strip().isin({"", "nan"})
        & ~y_pred.astype(str).str.strip().isin({"", "nan"})
    )
    yt = y_true[mask].astype(str).tolist()
    yp = y_pred[mask].astype(str).tolist()

    if len(yt) == 0:
        print(f"[{etiqueta}] Sin datos suficientes")
        return None, None, None

    acc = accuracy_score(yt, yp)
    p, r, f, _ = precision_recall_fscore_support(
        yt, yp, labels=LABELS_CAT, average="macro", zero_division=0
    )
    kappa = cohen_kappa_score(yt, yp, labels=LABELS_CAT)
    cm = confusion_matrix(yt, yp, labels=LABELS_CAT)
    n_correctos = int(sum(a == b for a, b in zip(yt, yp)))

    resumen = {
        "comparacion": etiqueta,
        "n_departamentos": len(yt),
        "accuracy": round(float(acc), 4),
        "n_correctos": n_correctos,
        "precision_macro": round(float(p), 4),
        "recall_macro": round(float(r), 4),
        "f1_macro": round(float(f), 4),
        "cohen_kappa": round(float(kappa), 4),
    }

    rep = classification_report(
        yt, yp, labels=LABELS_CAT, output_dict=True, zero_division=0
    )
    detalle_rows = []
    for clase in LABELS_CAT:
        d = rep[clase]
        detalle_rows.append({
            "comparacion": etiqueta,
            "clase": clase,
            "precision": round(float(d["precision"]), 4),
            "recall": round(float(d["recall"]), 4),
            "f1-score": round(float(d["f1-score"]), 4),
            "support": int(d["support"]),
        })
    detalle_df = pd.DataFrame(detalle_rows)

    return resumen, detalle_df, cm


# ---------------------------------------------------------------------------
# Leer referencia oficial V3
# ---------------------------------------------------------------------------

def _leer_oficial_v3(ruta_v3: str) -> pd.DataFrame:
    """Lee columnas oficiales DANE de comparacion_radares_V3.xlsx."""
    df = pd.read_excel(ruta_v3, engine="openpyxl")
    df.columns = [str(c).strip() for c in df.columns]
    cols_map = {c.casefold(): c for c in df.columns}
    dep = cols_map.get("departamento")
    ofi = cols_map.get("radar_oficial_promedio")
    nor = cols_map.get("radar_oficial_promedio_normalizado")
    cla = cols_map.get("clasificacion_radar_oficial_promedio")
    if not all([dep, ofi, nor, cla]):
        raise ValueError(
            f"'{ruta_v3}' debe tener: Departamento, radar_oficial_promedio, "
            "radar_oficial_promedio_normalizado, Clasificacion_radar_oficial_promedio"
        )
    out = df[[dep, ofi, nor, cla]].copy()
    out.columns = ["Departamento", "radar_oficial_promedio",
                   "radar_oficial_promedio_normalizado", "Clasificacion_radar_oficial_promedio"]
    out["Departamento"] = out["Departamento"].astype(str).str.strip()
    return out


# ---------------------------------------------------------------------------
# Generar comparacion por experimento (formato V3)
# ---------------------------------------------------------------------------

def generar_comparacion_experimento(
    experimento_id: str,
    df_radar_exp: pd.DataFrame,
    dir_salida: str,
    ruta_v3: str,
) -> str:
    """
    Genera comparacion_EXPERIMENTO_XXX.xlsx en formato V3:
    Departamento | radar_oficial_promedio | radar_oficial_promedio_normalizado |
    Clasificacion_radar_oficial_promedio | {EXP}_normalizado | Clasificacion_{EXP}

    Incluye los 32 departamentos DANE; NaN para los que no tiene el experimento.
    """
    os.makedirs(dir_salida, exist_ok=True)
    df_oficial = _leer_oficial_v3(ruta_v3)

    df_exp = df_radar_exp[["departamento", "radar_propio"]].copy()
    df_exp["departamento"] = df_exp["departamento"].astype(str).str.strip()
    df_exp["radar_propio"] = pd.to_numeric(df_exp["radar_propio"], errors="coerce")

    col_norm = f"{experimento_id}_normalizado"
    col_clas = f"Clasificacion_{experimento_id}"
    df_exp[col_norm] = _normalizar_minmax(df_exp["radar_propio"])
    df_exp[col_clas] = _clasificar(df_exp["radar_propio"]).values

    df_exp["_dep_lower"] = df_exp["departamento"].str.casefold()
    df_oficial["_dep_lower"] = df_oficial["Departamento"].str.casefold()

    df_merged = df_oficial.merge(
        df_exp[["_dep_lower", col_norm, col_clas]],
        on="_dep_lower", how="left"
    ).drop(columns=["_dep_lower"])

    ruta_salida = os.path.join(dir_salida, f"comparacion_{experimento_id}.xlsx")
    df_merged.to_excel(ruta_salida, index=False, engine="openpyxl")
    return ruta_salida


# ---------------------------------------------------------------------------
# Actualizar metricas_experimentos.xlsx (acumulado)
# ---------------------------------------------------------------------------

def _actualizar_metricas_acumulado(
    comparacion_label: str,
    resumen: dict,
    detalle_df: pd.DataFrame,
    cm: np.ndarray,
    ruta_metricas: str,
) -> None:
    """Append/update en metricas_experimentos.xlsx, reemplazando si ya existe."""
    if os.path.isfile(ruta_metricas):
        sheets = pd.read_excel(ruta_metricas, sheet_name=None, engine="openpyxl")
        df_res = sheets.get("resumen", pd.DataFrame())
        df_det = sheets.get("detalle_por_clase", pd.DataFrame())
        df_cm_prev = sheets.get("matrices_confusion", pd.DataFrame())
    else:
        df_res = pd.DataFrame()
        df_det = pd.DataFrame()
        df_cm_prev = pd.DataFrame()

    # Eliminar entradas anteriores de este experimento
    if not df_res.empty and "comparacion" in df_res.columns:
        df_res = df_res[df_res["comparacion"] != comparacion_label].reset_index(drop=True)
    if not df_det.empty and "comparacion" in df_det.columns:
        df_det = df_det[df_det["comparacion"] != comparacion_label].reset_index(drop=True)
    if not df_cm_prev.empty and "comparacion" in df_cm_prev.columns:
        df_cm_prev = df_cm_prev[
            df_cm_prev["comparacion"].astype(str).str.strip().ne(comparacion_label)
            | df_cm_prev["comparacion"].isna()
        ]
        df_cm_prev = df_cm_prev.reset_index(drop=True)

    # Resumen
    cols_resumen = ["comparacion", "n_departamentos", "accuracy",
                    "precision_macro", "recall_macro", "f1_macro", "cohen_kappa",
                    "baseline_azar", "baseline_clase_mayoritaria",
                    "baseline_clase_mayoritaria_cual", "baseline_modelo_nulo_articulos"]
    nueva_fila = {k: resumen.get(k) for k in cols_resumen}
    df_res_new = pd.DataFrame([nueva_fila])
    df_res = pd.concat([df_res, df_res_new], ignore_index=True)
    if "comparacion" in df_res.columns:
        cols_ord = [c for c in cols_resumen if c in df_res.columns]
        df_res = df_res[cols_ord]

    # Detalle por clase
    df_det_cols = ["comparacion", "clase", "precision", "recall", "f1-score", "support"]
    det_sub = detalle_df.rename(columns={"f1_score": "f1-score"}) if "f1_score" in detalle_df.columns else detalle_df
    det_sub = det_sub[[c for c in df_det_cols if c in det_sub.columns]]
    df_det = pd.concat([df_det, det_sub], ignore_index=True)

    # Matrices de confusión
    filas_cm: List[dict] = []
    filas_cm.append({"comparacion": comparacion_label, "verdadero \\ predicho": "---",
                     **{c: "" for c in LABELS_CAT}})
    for ri, label_r in enumerate(LABELS_CAT):
        fila = {"comparacion": comparacion_label, "verdadero \\ predicho": label_r}
        for ci, label_c in enumerate(LABELS_CAT):
            fila[label_c] = int(cm[ri, ci])
        filas_cm.append(fila)
    filas_cm.append({"comparacion": "", "verdadero \\ predicho": ""})
    df_cm_new = pd.DataFrame(filas_cm)

    if df_cm_prev.empty:
        df_cm_out = df_cm_new
    else:
        df_cm_out = pd.concat([df_cm_prev, df_cm_new], ignore_index=True)

    with pd.ExcelWriter(ruta_metricas, engine="openpyxl") as writer:
        _reemplazar_inf_para_excel(df_res).to_excel(writer, sheet_name="resumen", index=False)
        _reemplazar_inf_para_excel(df_det).to_excel(writer, sheet_name="detalle_por_clase", index=False)
        df_cm_out.to_excel(writer, sheet_name="matrices_confusion", index=False)

    print(f"[metricas] Acumulado actualizado: {ruta_metricas}")


# ---------------------------------------------------------------------------
# Generar gráficos por experimento
# ---------------------------------------------------------------------------

def _generar_graficos_experimento(
    experimento_id: str,
    resumen: dict,
    cm: np.ndarray,
    df_joined: pd.DataFrame,
    cat_oficial: pd.Series,
    cat_exp: pd.Series,
    dir_salida: str,
) -> None:
    os.makedirs(dir_salida, exist_ok=True)
    comparacion_label = resumen["comparacion"]

    cm_df = pd.DataFrame(cm, index=LABELS_CAT, columns=LABELS_CAT)
    plt.figure(figsize=(6, 5))
    sns.heatmap(cm_df, annot=True, fmt="d", cmap="Blues")
    plt.title(
        f"{experimento_id}\n"
        f"Acc={resumen['accuracy']*100:.1f}%  F1={resumen['f1_macro']:.3f}  "
        f"kappa={resumen['cohen_kappa']:.3f}  "
        f"({resumen['n_correctos']}/{resumen['n_departamentos']})"
    )
    plt.xlabel("Prediccion (experimento)")
    plt.ylabel("Real (oficial DANE)")
    plt.tight_layout()
    plt.savefig(os.path.join(dir_salida, "confusion_matrix.png"), dpi=150)
    plt.close()

    ts = df_joined.copy()
    ts["categoria_oficial"] = cat_oficial.values
    ts["categoria_experimento"] = cat_exp.values
    ts["correcto"] = (cat_oficial.values == cat_exp.values)
    ts = ts.sort_values("departamento").reset_index(drop=True)
    colores = ts["correcto"].map({True: "steelblue", False: "tomato"})
    fig, ax = plt.subplots(figsize=(14, 5))
    ax.scatter(ts.index, [1] * len(ts), color=colores, s=140, zorder=3)
    for i, row in ts.iterrows():
        dep = str(row.get("departamento", ""))[:10]
        cat_of = row["categoria_oficial"]
        cat_ex = row["categoria_experimento"]
        ax.annotate(
            f"{dep}\nOf:{cat_of}\nEx:{cat_ex}",
            (i, 1), textcoords="offset points", xytext=(0, -48), fontsize=6, ha="center",
        )
    ax.set_xlim(-1, len(ts))
    ax.set_ylim(0.4, 1.5)
    ax.axis("off")
    ax.set_title(f"{experimento_id}: azul = correcto, rojo = error")
    plt.tight_layout()
    plt.savefig(os.path.join(dir_salida, "categorias_departamentos.png"), dpi=150)
    plt.close()


# ---------------------------------------------------------------------------
# Entrada principal — por experimento
# ---------------------------------------------------------------------------

def calcular_metricas_experimento(
    experimento_id: str,
    ruta_radar_csv: str,
    ruta_v3: str,
    dir_experimento: str,
    ruta_metricas_acumulado: str,
) -> Dict[str, Any]:
    """
    Para un experimento:
    1. Genera comparacion_{EXP}.xlsx (formato V3) en dir_experimento
    2. Calcula métricas vs radar oficial DANE (intersección de departamentos)
    3. Actualiza metricas_experimentos.xlsx (acumulado)
    4. Genera gráficos en dir_experimento
    """
    os.makedirs(dir_experimento, exist_ok=True)

    df_radar = pd.read_csv(ruta_radar_csv)
    df_radar["departamento"] = df_radar["departamento"].astype(str).str.strip()
    df_radar["radar_propio"] = pd.to_numeric(df_radar["radar_propio"], errors="coerce")
    if "n_articulos" in df_radar.columns:
        df_radar["n_articulos"] = pd.to_numeric(df_radar["n_articulos"], errors="coerce")

    df_oficial = _leer_oficial_v3(ruta_v3)

    # 1. Generar comparacion Excel (32 DANE departments, NaN para los que faltan)
    ruta_comp = generar_comparacion_experimento(experimento_id, df_radar, dir_experimento, ruta_v3)
    print(f"[metricas] Comparacion: {ruta_comp}")

    # 2. Calcular metricas sobre interseccion
    df_oficial["_dep_lower"] = df_oficial["Departamento"].str.casefold()
    df_radar["_dep_lower"] = df_radar["departamento"].str.casefold()
    cols_radar = ["_dep_lower", "departamento", "radar_propio"]
    if "n_articulos" in df_radar.columns:
        cols_radar.append("n_articulos")
    df_joined = df_oficial.merge(
        df_radar[cols_radar],
        on="_dep_lower", how="inner"
    ).drop(columns=["_dep_lower"])

    if len(df_joined) < 3:
        raise ValueError(
            f"Interseccion insuficiente ({len(df_joined)} departamentos) "
            "entre radar oficial y experimento"
        )

    cat_oficial = _clasificar(_to_numeric_series(df_joined["radar_oficial_promedio"]))
    cat_exp = _clasificar(df_joined["radar_propio"])

    comparacion_label = f"{experimento_id}_vs_radar_oficial"
    resumen, detalle_df, cm = _metricas_par(cat_oficial, cat_exp, comparacion_label)
    if resumen is None:
        raise ValueError(f"No se pudieron calcular metricas para {experimento_id}")

    n_articulos_joined = df_joined["n_articulos"] if "n_articulos" in df_joined.columns else None
    base = _lineas_base(cat_oficial, n_articulos_joined)
    resumen.update(base)

    # Consola
    print(f"\n========== METRICAS: {comparacion_label} ==========")
    print(f"  n={resumen['n_departamentos']}  "
          f"accuracy={resumen['accuracy']*100:.1f}%  "
          f"F1={resumen['f1_macro']:.3f}  "
          f"kappa={resumen['cohen_kappa']:.3f}  "
          f"({resumen['n_correctos']}/{resumen['n_departamentos']})")
    _imprimir_lineas_base(base, resumen["accuracy"])

    errores = df_joined.copy()
    errores["cat_oficial"] = cat_oficial.values
    errores["cat_exp"] = cat_exp.values
    errores = errores[errores["cat_oficial"] != errores["cat_exp"]]
    if not errores.empty:
        print("  Errores de clasificacion:")
        for _, row in errores.iterrows():
            print(f"    {row['departamento']}: oficial={row['cat_oficial']}  exp={row['cat_exp']}")

    # 3. Actualizar acumulado
    _actualizar_metricas_acumulado(comparacion_label, resumen, detalle_df, cm, ruta_metricas_acumulado)

    # 4. Graficos
    _generar_graficos_experimento(
        experimento_id, resumen, cm, df_joined, cat_oficial, cat_exp, dir_experimento
    )

    return {
        "metricas": [resumen],
        "best_accuracy": resumen["accuracy"],
        "dir_experimento": dir_experimento,
        "excel_comparacion": ruta_comp,
        "excel_metricas": ruta_metricas_acumulado,
        "experimentos_procesados": [comparacion_label],
        "experimentos_sin_datos": [],
    }


# ---------------------------------------------------------------------------
# Legado — conservado para compatibilidad con llamadas directas antiguas
# ---------------------------------------------------------------------------

def _asegurar_clasificacion_oficial(ruta_excel: str) -> None:
    df = pd.read_excel(ruta_excel, engine="openpyxl")
    if COL_CLASIFICACION_OFICIAL in df.columns:
        return
    if "radar_oficial_promedio" not in df.columns:
        raise ValueError(f"'{ruta_excel}' no tiene columna 'radar_oficial_promedio'")
    serie_oficial = _to_numeric_series(df["radar_oficial_promedio"])
    clasificacion = _clasificar(serie_oficial)
    cols = df.columns.tolist()
    idx = cols.index("radar_oficial_promedio") + 1
    df.insert(idx, COL_CLASIFICACION_OFICIAL, clasificacion)
    df.to_excel(ruta_excel, index=False, engine="openpyxl")
    print(f"[metricas] Columna '{COL_CLASIFICACION_OFICIAL}' agregada a '{ruta_excel}'")
    print(f"  Distribucion oficial: {clasificacion.value_counts().to_dict()}")


def _cargar_excel_con_headers_reales(ruta_excel: str) -> pd.DataFrame:
    wb = load_workbook(ruta_excel, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return pd.DataFrame()
    headers = ["" if c is None else str(c).strip() for c in filas[0]]
    max_header = max(
        (idx for idx, h in enumerate(headers, 1) if h and not h.lower().startswith("unnamed:")),
        default=0,
    )
    if max_header == 0:
        return pd.DataFrame()
    headers = headers[:max_header]
    data_rows = []
    for fila in filas[1:]:
        if not fila:
            continue
        recorte = list(fila[:max_header])
        recorte += [None] * (max_header - len(recorte))
        if all(c is None for c in recorte):
            continue
        data_rows.append(recorte)
    return pd.DataFrame(data_rows, columns=headers)


def procesar_metricas_multi_experimento(ruta_excel: str, salida: str) -> Dict[str, Any]:
    """LEGADO: compara múltiples experimentos desde el Excel ancho. Usar calcular_metricas_experimento()."""
    os.makedirs(salida, exist_ok=True)
    graficos_dir = os.path.join(salida, cfg.DIRECTORIO_GRAFICOS_METRICAS)
    excel_salida = os.path.join(salida, f"{cfg.PREFIJO_RESULTADO_COMPARACION}.xlsx")

    _asegurar_clasificacion_oficial(ruta_excel)
    df_raw = _cargar_excel_con_headers_reales(ruta_excel)
    if df_raw.shape[1] == 0:
        raise ValueError("El archivo de comparacion esta vacio")

    contador: Dict[str, int] = {}
    headers_internos: List[str] = []
    header_orig: Dict[str, str] = {}
    for h in (str(c).strip() if c else "" for c in df_raw.columns):
        clave = h.casefold()
        contador[clave] = contador.get(clave, 0) + 1
        interno = h if contador[clave] == 1 else f"{h}__dup_{contador[clave]}"
        headers_internos.append(interno)
        header_orig[interno] = h

    df = df_raw.copy()
    df.columns = headers_internos

    col_dep = next((c for c in df.columns if header_orig[c].casefold() == "departamento"), None)
    col_oficial = next((c for c in df.columns if header_orig[c].casefold() == "radar_oficial_promedio"), None)
    if not col_dep or not col_oficial:
        raise ValueError("Faltan columnas obligatorias en el Excel")

    col_clasif_oficial = next(
        (c for c in df.columns if header_orig[c].casefold() == COL_CLASIFICACION_OFICIAL.casefold()), None
    )
    cols_excluir = {col_dep, col_oficial}
    if col_clasif_oficial:
        cols_excluir.add(col_clasif_oficial)
    experimentos = [c for c in df.columns if c not in cols_excluir
                    and header_orig[c] and not header_orig[c].lower().startswith("unnamed:")]

    if not experimentos:
        raise ValueError("No se encontraron columnas de experimento")

    df[col_dep] = df[col_dep].astype(str).str.strip()
    df[col_oficial] = _to_numeric_series(df[col_oficial])
    for col in experimentos:
        df[col] = _to_numeric_series(df[col])

    dist_oficial = _clasificar(df[col_oficial]).value_counts().to_dict()
    print(f"\nDistribucion oficial: {dist_oficial}")

    ids_exp: Dict[str, str] = {}
    contador_ids: Dict[str, int] = {}
    carpetas_exp: Dict[str, str] = {}
    carpetas_usadas: Dict[str, int] = {}
    os.makedirs(graficos_dir, exist_ok=True)
    for col in experimentos:
        h = header_orig[col]
        clave = h.casefold()
        contador_ids[clave] = contador_ids.get(clave, 0) + 1
        eid = h if contador_ids[clave] == 1 else f"{h}__{contador_ids[clave]}"
        ids_exp[col] = eid
        nc = _nombre_carpeta_seguro(eid)
        ck = nc.casefold()
        carpetas_usadas[ck] = carpetas_usadas.get(ck, 0) + 1
        if carpetas_usadas[ck] > 1:
            nc = f"{nc}__{carpetas_usadas[ck]}"
        carpetas_exp[eid] = os.path.join(graficos_dir, nc)

    resumenes: List[dict] = []
    detalles: List[pd.DataFrame] = []
    clasificaciones: List[pd.DataFrame] = []
    sin_datos: List[str] = []

    for col in experimentos:
        eid = ids_exp[col]
        temp = df[[col_dep, col_oficial, col]].copy()
        temp.columns = ["departamento", "radar_oficial_promedio", "radar_experimento"]
        temp = temp.dropna(subset=["radar_oficial_promedio", "radar_experimento"])
        if len(temp) < 3:
            sin_datos.append(eid)
            continue

        cat_of = _clasificar(temp["radar_oficial_promedio"])
        cat_ex = _clasificar(temp["radar_experimento"])
        res, det, cm = _metricas_par(cat_of, cat_ex, eid)
        if res is None:
            sin_datos.append(eid)
            continue

        # Adapt key names for legacy compatibility
        res_legacy = {k.replace("comparacion", "experimento") if k == "comparacion" else k: v
                      for k, v in res.items()}
        res_legacy["experimento"] = eid

        resumenes.append(res_legacy)
        det_legacy = det.rename(columns={"comparacion": "experimento"}) if det is not None else det
        detalles.append(det_legacy)

        temp["categoria_oficial"] = cat_of.values
        temp["categoria_experimento"] = cat_ex.values
        temp["correcto"] = (cat_of == cat_ex).values
        temp["experimento"] = eid
        clasificaciones.append(
            temp[["experimento", "departamento", "radar_oficial_promedio", "radar_experimento",
                  "categoria_oficial", "categoria_experimento", "correcto"]]
        )
        subdir = carpetas_exp[eid]
        os.makedirs(subdir, exist_ok=True)
        cm_df = pd.DataFrame(cm, index=LABELS_CAT, columns=LABELS_CAT)
        plt.figure(figsize=(6, 5))
        sns.heatmap(cm_df, annot=True, fmt="d", cmap="Blues")
        plt.title(f"{eid}\nAcc={res_legacy['accuracy']*100:.1f}%")
        plt.tight_layout()
        plt.savefig(os.path.join(subdir, "confusion_matrix.png"), dpi=150)
        plt.close()

    if not resumenes:
        raise ValueError("No hay datos suficientes para calcular metricas")

    df_res = pd.DataFrame(resumenes).sort_values("accuracy", ascending=False).reset_index(drop=True)
    for col_m, _ in [("accuracy", 0.50), ("f1_macro", 0.30), ("cohen_kappa", 0.20)]:
        df_res[f"rank_{col_m}"] = df_res[col_m].rank(method="min", ascending=False)
    df_res["score_ranking"] = (
        0.50 * df_res["rank_accuracy"]
        + 0.30 * df_res["rank_f1_macro"]
        + 0.20 * df_res["rank_cohen_kappa"]
    )

    df_det = pd.concat(detalles, ignore_index=True) if detalles else pd.DataFrame()
    df_clasi = pd.concat(clasificaciones, ignore_index=True) if clasificaciones else pd.DataFrame()

    filas_cm: List[dict] = []
    for col in experimentos:
        eid = ids_exp[col]
        if eid not in [r.get("experimento") for r in resumenes]:
            continue
        temp_c = df[[col_dep, col_oficial, col]].dropna(subset=[col_oficial, col]).copy()
        cat_of_c = _clasificar(_to_numeric_series(temp_c[col_oficial]))
        cat_ex_c = _clasificar(_to_numeric_series(temp_c[col]))
        cm_c = confusion_matrix(cat_of_c.tolist(), cat_ex_c.tolist(), labels=LABELS_CAT)
        filas_cm.append({"experimento": eid, "real \\ predicho": "---", **{c: "" for c in LABELS_CAT}})
        for ri, label_r in enumerate(LABELS_CAT):
            fila = {"experimento": eid, "real \\ predicho": label_r}
            for ci, label_c in enumerate(LABELS_CAT):
                fila[label_c] = int(cm_c[ri, ci])
            filas_cm.append(fila)
        filas_cm.append({"experimento": "", "real \\ predicho": ""})
    df_cm_excel = pd.DataFrame(filas_cm)

    print("\n========== METRICAS DE CLASIFICACION ==========")
    cols_print = ["experimento", "n_departamentos", "accuracy", "n_correctos", "f1_macro", "cohen_kappa"]
    cols_print = [c for c in cols_print if c in df_res.columns]
    print(df_res[cols_print].to_string(index=False))

    plt.figure(figsize=(max(8, len(df_res) * 0.7), 5))
    colores_bar = ["steelblue" if a >= 0.60 else "tomato" for a in df_res["accuracy"]]
    sns.barplot(data=df_res, x="experimento", y="accuracy", palette=colores_bar)
    plt.axhline(0.60, color="red", linestyle="--", alpha=0.7, label="60%")
    plt.axhline(0.70, color="green", linestyle="--", alpha=0.7, label="70%")
    plt.title("Accuracy por experimento")
    plt.ylabel("Accuracy")
    plt.ylim(0, 1)
    plt.xticks(rotation=45, ha="right")
    plt.legend()
    plt.tight_layout()
    plt.savefig(os.path.join(graficos_dir, "accuracy_experimentos.png"), dpi=150)
    plt.close()

    with pd.ExcelWriter(excel_salida, engine="openpyxl") as writer:
        _reemplazar_inf_para_excel(df_res).to_excel(writer, sheet_name="resumen", index=False)
        _reemplazar_inf_para_excel(df_det).to_excel(writer, sheet_name="detalle_por_clase", index=False)
        df_cm_excel.to_excel(writer, sheet_name="matrices_confusion", index=False)
        df_clasi.to_excel(writer, sheet_name="clasificacion_departamentos", index=False)

    print(f"\nArchivo exportado: {excel_salida}")
    best_exp = df_res.iloc[0]["experimento"] if not df_res.empty else ""
    best_metrics = {k: df_res.iloc[0][k] for k in ["accuracy", "f1_macro", "cohen_kappa"]
                    if k in df_res.columns} if not df_res.empty else {}

    return {
        "excel_salida": excel_salida,
        "graficos_dir": graficos_dir,
        "metricas": df_res.rename(columns={"experimento": "comparacion"}).to_dict(orient="records"),
        "best_accuracy": float(df_res["accuracy"].max()),
        "experimentos_procesados": df_res["experimento"].tolist(),
        "experimentos_sin_datos": sin_datos,
    }


def calcular_metricas(ruta_excel: str, salida: str) -> Dict[str, Any]:
    """LEGADO: wrapper hacia procesar_metricas_multi_experimento."""
    return procesar_metricas_multi_experimento(ruta_excel, salida)


if __name__ == "__main__":
    procesar_metricas_multi_experimento(cfg.ARCHIVO_COMPARACION_EXCEL, ".")
