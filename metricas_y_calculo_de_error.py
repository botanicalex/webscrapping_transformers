# -*- coding: utf-8 -*-
import os
from typing import Any, Dict, List
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import load_workbook
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from scipy.stats import pearsonr

import config_pipeline as cfg


def _to_numeric_series(serie: pd.Series) -> pd.Series:
    serie_limpia = serie.copy()
    tokens_nulos = {"none", "null", "nan", ""}
    mascara_nulos = serie_limpia.astype(str).str.strip().str.casefold().isin(tokens_nulos)
    serie_limpia = serie_limpia.mask(mascara_nulos, np.nan)
    serie_limpia = serie_limpia.astype(str).str.replace(",", ".", regex=False).str.strip()
    return pd.to_numeric(serie_limpia, errors="coerce")


def _safe_pearson(y_real: pd.Series, y_pred: pd.Series) -> tuple[float, float]:
    if y_real.nunique() < 2 or y_pred.nunique() < 2:
        return np.nan, np.nan
    try:
        corr, pvalue = pearsonr(y_real, y_pred)
        return float(corr), float(pvalue)
    except Exception:
        return np.nan, np.nan


def _nombre_carpeta_seguro(nombre: str) -> str:
    nombre_base = str(nombre).strip()
    if not nombre_base:
        nombre_base = "experimento"
    nombre_limpio = "".join(
        c if c.isalnum() or c in (" ", "-", "_") else "_" for c in nombre_base
    ).strip()
    if not nombre_limpio:
        return "experimento"
    return nombre_limpio[:120]


def _reemplazar_inf_para_excel(df: pd.DataFrame) -> pd.DataFrame:
    return df.replace([np.inf, -np.inf], np.nan)


def _cargar_excel_con_headers_reales(ruta_excel: str) -> pd.DataFrame:
    wb = load_workbook(ruta_excel, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return pd.DataFrame()

    headers = ["" if c is None else str(c).strip() for c in filas[0]]
    max_header = 0
    for idx, h in enumerate(headers, start=1):
        if h != "" and not h.lower().startswith("unnamed:"):
            max_header = idx

    if max_header == 0:
        return pd.DataFrame()

    headers = headers[:max_header]
    data_rows = []
    for fila in filas[1:]:
        if fila is None:
            continue
        recorte = list(fila[:max_header])
        if len(recorte) < max_header:
            recorte += [None] * (max_header - len(recorte))
        if all(celda is None for celda in recorte):
            continue
        data_rows.append(recorte)

    return pd.DataFrame(data_rows, columns=headers)


def procesar_metricas_multi_experimento(ruta_excel: str, salida: str) -> Dict[str, Any]:
    os.makedirs(salida, exist_ok=True)
    graficos_dir = os.path.join(salida, cfg.DIRECTORIO_GRAFICOS_METRICAS)
    excel_salida = os.path.join(salida, f"{cfg.PREFIJO_RESULTADO_COMPARACION}.xlsx")

    df_raw = _cargar_excel_con_headers_reales(ruta_excel)
    if df_raw.shape[1] == 0:
        raise ValueError("El archivo de comparación está vacío o no contiene columnas")
    if df_raw.shape[0] == 0:
        raise ValueError("El archivo de comparación no contiene filas de datos")

    headers_originales = ["" if c is None else str(c).strip() for c in df_raw.columns]
    contador_headers: Dict[str, int] = {}
    headers_internos: List[str] = []
    header_original_por_columna: Dict[str, str] = {}

    for header in headers_originales:
        clave = header.casefold()
        contador_headers[clave] = contador_headers.get(clave, 0) + 1
        if contador_headers[clave] == 1:
            header_interno = header
        else:
            header_interno = f"{header}__dup_{contador_headers[clave]}"
        headers_internos.append(header_interno)
        header_original_por_columna[header_interno] = header

    df = df_raw.copy()
    df.columns = headers_internos

    columnas_departamento = [
        c for c in df.columns if header_original_por_columna[c].casefold() == "departamento"
    ]
    if not columnas_departamento:
        raise ValueError("Falta la columna requerida: departamento")
    if len(columnas_departamento) > 1:
        raise ValueError("Hay columnas duplicadas para 'departamento' y no se puede resolver la ambigüedad")

    columnas_radar_oficial = [
        c for c in df.columns if header_original_por_columna[c].casefold() == "radar_oficial_promedio"
    ]
    if not columnas_radar_oficial:
        raise ValueError("Falta la columna requerida: radar_oficial_promedio")
    if len(columnas_radar_oficial) > 1:
        raise ValueError(
            "Hay columnas duplicadas para 'radar_oficial_promedio' y no se puede resolver la ambigüedad"
        )

    col_departamento = columnas_departamento[0]
    col_radar_oficial = columnas_radar_oficial[0]

    experimentos: List[str] = []
    for c in df.columns:
        if c in (col_departamento, col_radar_oficial):
            continue
        header_original = header_original_por_columna[c]
        if header_original == "" or header_original.lower().startswith("unnamed:"):
            continue
        experimentos.append(c)

    if not experimentos:
        raise ValueError("No se encontraron columnas de experimento en el archivo de comparación")

    ids_experimento: Dict[str, str] = {}
    contador_ids: Dict[str, int] = {}
    carpetas_usadas: Dict[str, int] = {}
    carpetas_experimento: Dict[str, str] = {}

    for exp_col in experimentos:
        header_original = header_original_por_columna[exp_col]
        clave = header_original.casefold()
        contador_ids[clave] = contador_ids.get(clave, 0) + 1
        if contador_ids[clave] == 1:
            experimento_id = header_original
        else:
            experimento_id = f"{header_original}__{contador_ids[clave]}"
        ids_experimento[exp_col] = experimento_id

        nombre_carpeta = _nombre_carpeta_seguro(experimento_id)
        clave_carpeta = nombre_carpeta.casefold()
        carpetas_usadas[clave_carpeta] = carpetas_usadas.get(clave_carpeta, 0) + 1
        if carpetas_usadas[clave_carpeta] > 1:
            nombre_carpeta = f"{nombre_carpeta}__{carpetas_usadas[clave_carpeta]}"
        carpetas_experimento[experimento_id] = os.path.join(graficos_dir, nombre_carpeta)

    df[col_departamento] = df[col_departamento].astype(str).str.strip()
    df[col_radar_oficial] = _to_numeric_series(df[col_radar_oficial])
    for col in experimentos:
        df[col] = _to_numeric_series(df[col])

    res_metricas: List[Dict[str, Any]] = []
    res_errores: List[pd.DataFrame] = []
    experimentos_sin_datos: List[str] = []

    os.makedirs(graficos_dir, exist_ok=True)

    for exp_col in experimentos:
        experimento_id = ids_experimento[exp_col]
        header_experimento = header_original_por_columna[exp_col]
        temp = df[[col_departamento, col_radar_oficial, exp_col]].copy()
        temp = temp.rename(
            columns={
                col_departamento: "departamento",
                col_radar_oficial: "radar_oficial_promedio",
                exp_col: "radar_experimento",
            }
        )
        temp = temp.dropna(subset=["radar_oficial_promedio", "radar_experimento"])

        if len(temp) < 2:
            experimentos_sin_datos.append(experimento_id)
            continue

        y_real = temp["radar_oficial_promedio"]
        y_pred = temp["radar_experimento"]

        mae = mean_absolute_error(y_real, y_pred)
        rmse = np.sqrt(mean_squared_error(y_real, y_pred))
        r2 = r2_score(y_real, y_pred)
        pearson_corr, pearson_p = _safe_pearson(y_real, y_pred)

        base_mape = y_real.replace(0, np.nan)
        mape = np.nanmean(np.abs((y_real - y_pred) / base_mape)) * 100
        sesgo = (y_pred - y_real).mean()
        if np.isnan(mape):
            mape = float("inf")

        res_metricas.append(
            {
                "experimento": experimento_id,
                "header_experimento": header_experimento,
                "n_departamentos": len(temp),
                "MAE": float(mae),
                "RMSE": float(rmse),
                "R2": float(r2),
                "Pearson": float(pearson_corr) if not np.isnan(pearson_corr) else np.nan,
                "Pearson_p": float(pearson_p) if not np.isnan(pearson_p) else np.nan,
                "MAPE_pct": float(mape),
                "sesgo_promedio": float(sesgo),
            }
        )

        temp["experimento"] = experimento_id
        temp["header_experimento"] = header_experimento
        temp["error"] = temp["radar_experimento"] - temp["radar_oficial_promedio"]
        temp["error_absoluto"] = temp["error"].abs()

        res_errores.append(
            temp[
                [
                    "experimento",
                    "header_experimento",
                    "departamento",
                    "radar_experimento",
                    "radar_oficial_promedio",
                    "error",
                    "error_absoluto",
                ]
            ]
        )

        subdir_experimento = carpetas_experimento[experimento_id]
        os.makedirs(subdir_experimento, exist_ok=True)

        plt.figure(figsize=(8, 6))
        sns.scatterplot(
            data=temp,
            x="radar_oficial_promedio",
            y="radar_experimento",
            s=90,
            color="steelblue",
        )
        min_val = min(temp["radar_oficial_promedio"].min(), temp["radar_experimento"].min())
        max_val = max(temp["radar_oficial_promedio"].max(), temp["radar_experimento"].max())
        plt.plot([min_val, max_val], [min_val, max_val], linestyle="--", color="red")
        plt.title(f"{experimento_id}: oficial vs experimento")
        plt.xlabel("Radar oficial promedio")
        plt.ylabel("Radar experimento")
        plt.tight_layout()
        plt.savefig(os.path.join(subdir_experimento, "scatter.png"), dpi=150)
        plt.close()

        temp_ordenado = temp.sort_values("error", ascending=False)
        plt.figure(figsize=(14, 6))
        sns.barplot(data=temp_ordenado, x="departamento", y="error", color="steelblue")
        plt.axhline(0, color="black", linewidth=1)
        plt.title(f"{experimento_id}: error por departamento (experimento - oficial)")
        plt.xlabel("Departamento")
        plt.ylabel("Error")
        plt.xticks(rotation=90)
        plt.tight_layout()
        plt.savefig(os.path.join(subdir_experimento, "error_departamento.png"), dpi=150)
        plt.close()

    if not res_metricas:
        raise ValueError("No hay datos suficientes para calcular métricas en los experimentos")

    df_metricas = pd.DataFrame(res_metricas)
    df_metricas = df_metricas.sort_values("MAE").reset_index(drop=True)

    if res_errores:
        df_errores = pd.concat(res_errores, ignore_index=True)
    else:
        df_errores = pd.DataFrame(
            columns=[
                "experimento",
                "header_experimento",
                "departamento",
                "radar_experimento",
                "radar_oficial_promedio",
                "error",
                "error_absoluto",
            ]
        )

    rank_base = df_metricas.copy()
    rank_base["rank_MAE"] = rank_base["MAE"].rank(method="min", ascending=True)
    rank_base["rank_RMSE"] = rank_base["RMSE"].rank(method="min", ascending=True)
    rank_base["rank_MAPE"] = rank_base["MAPE_pct"].rank(method="min", ascending=True)
    rank_base["rank_Pearson"] = rank_base["Pearson"].fillna(float("-inf")).rank(
        method="min", ascending=False
    )
    rank_base["score_ranking"] = (
        0.35 * rank_base["rank_MAE"]
        + 0.30 * rank_base["rank_RMSE"]
        + 0.20 * rank_base["rank_MAPE"]
        + 0.15 * rank_base["rank_Pearson"]
    )
    df_ranking = rank_base.sort_values("score_ranking").reset_index(drop=True)

    print("\n========== MÉTRICAS POR EXPERIMENTO ==========")
    print(df_metricas.to_string(index=False))

    if not df_errores.empty:
        print("\n========== TOP 10 ERRORES ABSOLUTOS GLOBALES ==========")
        print(df_errores.sort_values("error_absoluto", ascending=False).head(10).to_string(index=False))

    print("\n========== RANKING GLOBAL DE EXPERIMENTOS ==========")
    print(
        df_ranking[["experimento", "score_ranking", "MAE", "RMSE", "MAPE_pct", "Pearson"]].to_string(
            index=False
        )
    )

    if experimentos_sin_datos:
        print("\n========== EXPERIMENTOS SIN DATOS SUFICIENTES ==========")
        print(pd.DataFrame({"experimento": experimentos_sin_datos}).to_string(index=False))

    plt.figure(figsize=(12, 6))
    plot_rank = df_ranking[["experimento", "score_ranking"]].copy()
    sns.barplot(data=plot_rank, x="experimento", y="score_ranking", color="steelblue")
    plt.title("Ranking global de experimentos (menor es mejor)")
    plt.xlabel("Experimento")
    plt.ylabel("Score ranking")
    plt.xticks(rotation=45, ha="right")
    plt.tight_layout()
    plt.savefig(os.path.join(graficos_dir, "ranking_experimentos.png"), dpi=150)
    plt.close()

    df_metricas_export = _reemplazar_inf_para_excel(df_metricas)
    df_ranking_export = _reemplazar_inf_para_excel(df_ranking)
    df_errores_export = _reemplazar_inf_para_excel(df_errores)

    with pd.ExcelWriter(excel_salida, engine="openpyxl") as writer:
        df_metricas_export.to_excel(writer, sheet_name="metricas", index=False)
        df_ranking_export.to_excel(writer, sheet_name="ranking", index=False)
        df_errores_export.to_excel(writer, sheet_name="errores", index=False)
        if experimentos_sin_datos:
            pd.DataFrame({"experimento": experimentos_sin_datos}).to_excel(
                writer, sheet_name="experimentos_sin_datos", index=False
            )

    print(f"\nArchivo exportado: {excel_salida}")
    print(f"Gráficos exportados en: {graficos_dir}")

    max_error_abs = float(df_errores["error_absoluto"].max()) if not df_errores.empty else 0.0

    return {
        "excel_salida": excel_salida,
        "graficos_dir": graficos_dir,
        "metricas": df_metricas.to_dict(orient="records"),
        "max_error_abs": max_error_abs,
        "experimentos_procesados": df_metricas["experimento"].tolist(),
        "experimentos_sin_datos": experimentos_sin_datos,
    }


def calcular_metricas(ruta_excel: str, salida: str) -> Dict[str, Any]:
    return procesar_metricas_multi_experimento(ruta_excel, salida)


if __name__ == "__main__":
    procesar_metricas_multi_experimento(cfg.ARCHIVO_COMPARACION_EXCEL, ".")
